import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============================== INPUTS / CONSTANTS ==============================
NETREV      = 83_000_000        # July total net revenue (ex-GST, post-returns) — SCALE ANCHOR (estimate)
NEW_SHARE   = 0.506             # new-user share of net revenue (July-31 DMR sample)
AOV_NEW     = 2602
AOV_REP     = 3111
NG          = 0.7477            # Net Rev / Gross Rev  (COGS file)
COGS_R      = 0.4483            # Net COGS / Net Rev
FUL_R       = 0.0914            # Fulfilment / Net Rev
CM1_R       = 1 - COGS_R - FUL_R  # 0.4603
CRM         = 500_000

# Marketing actuals (July)
META        = 26_659_664
GOOGLE      = 8_641_852
META_PROSP  = 16_956_520
META_RETARG = 9_703_144
G_PMAX      = 8_152_672
G_BRAND     = 489_181

# GA4 observed paid mix (July, India acct)
META_NEW_GA4 = 0.453
GOOG_NEW_GA4 = 0.569

# ---- segment revenue / margins (attribution-independent) ----
def seg(nr, aov):
    gross = nr / NG
    return dict(nr=nr, gross=gross, orders=gross/aov,
                cogs=nr*COGS_R, pm=nr*(1-COGS_R), ful=nr*FUL_R, cm1=nr*CM1_R)
N = seg(NETREV*NEW_SHARE, AOV_NEW)
R = seg(NETREV*(1-NEW_SHARE), AOV_REP)

# ---- marketing allocation, two scenarios ----
# Intent-based
int_new = META_PROSP + META_RETARG*0.20 + G_PMAX*0.70 + G_BRAND*0.40
int_rep = META_RETARG*0.80 + G_PMAX*0.30 + G_BRAND*0.60 + CRM
# GA4 last-click
ga_new  = META*META_NEW_GA4 + GOOGLE*GOOG_NEW_GA4
ga_rep  = META*(1-META_NEW_GA4) + GOOGLE*(1-GOOG_NEW_GA4) + CRM

def cm2(cm1, media): return cm1 - media
scen = {
 "Intent": dict(nm=int_new, rm=int_rep,
                cn=cm2(N['cm1'], int_new), cr=cm2(R['cm1'], int_rep)),
 "GA4":    dict(nm=ga_new, rm=ga_rep,
                cn=cm2(N['cm1'], ga_new), cr=cm2(R['cm1'], ga_rep)),
}

# ============================== STYLES ==============================
H1=Font(bold=True,size=14,color="FFFFFF"); H2=Font(bold=True,size=12,color="1F3864")
BOLD=Font(bold=True); WHITE=Font(bold=True,color="FFFFFF"); ITAL=Font(italic=True,size=9,color="666666")
navy=PatternFill("solid",fgColor="1F3864"); blue=PatternFill("solid",fgColor="2E5496")
lblue=PatternFill("solid",fgColor="D9E1F2"); yell=PatternFill("solid",fgColor="FFF2CC")
grey=PatternFill("solid",fgColor="F2F2F2"); green=PatternFill("solid",fgColor="E2EFDA"); red=PatternFill("solid",fgColor="FCE4D6")
thin=Side(style="thin",color="BFBFBF"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
RS="#,##0"; PCT="0.0%"

def C(ws,addr,v,font=None,fill=None,fmt=None,al=None,bd=True):
    x=ws[addr]; x.value=v
    if font:x.font=font
    if fill:x.fill=fill
    if fmt:x.number_format=fmt
    if al: x.alignment=Alignment(horizontal=("left" if al=="wrap" else al),vertical="center",wrap_text=(al=="wrap"))
    if bd:x.border=box
    return x

wb=openpyxl.Workbook()

# ============================== P&L (values) ==============================
p=wb.active; p.title="P&L — New vs Repeat"; p.sheet_view.showGridLines=False
for c,w in {"A":2,"B":36,"C":15,"D":15,"E":15,"F":10}.items(): p.column_dimensions[c].width=w
p.merge_cells("B2:F2"); C(p,"B2","DailyObjects · July 2026 · D2C P&L to CM2 — New vs Repeat",H1,navy,al="left",bd=False)
p.merge_cells("B3:F3"); C(p,"B3","₹. Margins from COGS file; segment mix from July-31 DMR sample; media = July actuals. CM2 shown under BOTH attribution lenses.",ITAL,al="left",bd=False)

r=5
for c,t,a in [("B","P&L line","left"),("C","New user","center"),("D","Repeat user","center"),("E","Total","center"),("F","% NR","center")]:
    C(p,f"{c}{r}",t,WHITE,blue,al=a)
def row(r,label,nv,rv,fill=None,bold=False,pct=True,indent=True,money=True):
    f=BOLD if bold else None
    C(p,f"B{r}",("   " if indent else "")+label,f,fill,al="left")
    C(p,f"C{r}",round(nv),f,fill,RS if money else "#,##0",al="right")
    C(p,f"D{r}",round(rv),f,fill,RS if money else "#,##0",al="right")
    C(p,f"E{r}",round(nv+rv),f,fill,RS if money else "#,##0",al="right")
    C(p,f"F{r}",(nv+rv)/NETREV if pct else None,f,fill,PCT if pct else None,al="center")
r=6
row(r,"Net Revenue (ex-GST, post-returns)",N['nr'],R['nr'],lblue,True); r+=1
row(r,"Gross revenue (memo, incl GST+returns)",N['gross'],R['gross'],grey); r+=1
row(r,"Orders (derived = gross ÷ AOV)",N['orders'],R['orders'],grey,pct=False,money=False); r+=1
row(r,"Less: Net COGS (44.8%)",-N['cogs'],-R['cogs']); r+=1
row(r,"= Product Margin",N['pm'],R['pm'],green,True); r+=1
row(r,"Less: Fulfilment / logistics (9.1%)",-N['ful'],-R['ful']); r+=1
row(r,"= CM1 (post-fulfilment)",N['cm1'],R['cm1'],green,True); CM1ROW=r; r+=2

# --- CM2 under both scenarios ---
C(p,f"B{r}","CM2 — by marketing-attribution lens",WHITE,blue,al="left")
for c in "CDE": C(p,f"{c}{r}",{"C":"New","D":"Repeat","E":"Total"}[c],WHITE,blue,al="center")
C(p,f"F{r}","% NR",WHITE,blue,al="center"); r+=1
for key,nm_lbl,fill in [("Intent","Intent-based (Meta prospecting = New)",red),("GA4","GA4 last-click (observed)",green)]:
    s=scen[key]
    C(p,f"B{r}",f"   Marketing (Meta+Google+CRM) — {nm_lbl}",None,None,al="left")
    C(p,f"C{r}",-round(s['nm']),None,None,RS,al="right"); C(p,f"D{r}",-round(s['rm']),None,None,RS,al="right")
    C(p,f"E{r}",-round(s['nm']+s['rm']),None,None,RS,al="right"); C(p,f"F{r}",None); r+=1
    C(p,f"B{r}",f"   = CM2 — {key}",BOLD,fill,al="left")
    C(p,f"C{r}",round(s['cn']),BOLD,fill,RS,al="right"); C(p,f"D{r}",round(s['cr']),BOLD,fill,RS,al="right")
    C(p,f"E{r}",round(s['cn']+s['cr']),BOLD,fill,RS,al="right")
    C(p,f"F{r}",(s['cn']+s['cr'])/NETREV,BOLD,fill,PCT,al="center"); r+=1
    C(p,f"B{r}",f"      CM2 % of net revenue",ITAL,None,al="left")
    C(p,f"C{r}",s['cn']/N['nr'],ITAL,None,PCT,al="center"); C(p,f"D{r}",s['cr']/R['nr'],ITAL,None,PCT,al="center")
    C(p,f"E{r}",(s['cn']+s['cr'])/NETREV,ITAL,None,PCT,al="center"); C(p,f"F{r}",None); r+=1
r+=1

# --- unit economics memo ---
C(p,f"B{r}","MEMO — per order (₹)",WHITE,blue,al="left")
for c in "CDE": C(p,f"{c}{r}",{"C":"New","D":"Repeat","E":"Blended"}[c],WHITE,blue,al="center")
C(p,f"F{r}","",WHITE,blue); r+=1
tot_ord=N['orders']+R['orders']
C(p,f"B{r}","   Net revenue / order",None,None,al="left")
C(p,f"C{r}",round(N['nr']/N['orders']),None,None,RS,al="right");C(p,f"D{r}",round(R['nr']/R['orders']),None,None,RS,al="right");C(p,f"E{r}",round(NETREV/tot_ord),None,None,RS,al="right");C(p,f"F{r}","",bd=False); r+=1
C(p,f"B{r}","   CM1 / order",None,None,al="left")
C(p,f"C{r}",round(N['cm1']/N['orders']),None,None,RS,al="right");C(p,f"D{r}",round(R['cm1']/R['orders']),None,None,RS,al="right");C(p,f"E{r}",round((N['cm1']+R['cm1'])/tot_ord),None,None,RS,al="right");C(p,f"F{r}","",bd=False); r+=1
for key,lbl in [("Intent","Intent"),("GA4","GA4")]:
    s=scen[key]
    C(p,f"B{r}",f"   Marketing / order — {lbl}",None,None,al="left")
    C(p,f"C{r}",round(s['nm']/N['orders']),None,None,RS,al="right");C(p,f"D{r}",round(s['rm']/R['orders']),None,None,RS,al="right");C(p,f"E{r}",round((s['nm']+s['rm'])/tot_ord),None,None,RS,al="right");C(p,f"F{r}","",bd=False); r+=1
    C(p,f"B{r}",f"   CM2 / order — {lbl}",BOLD,None,al="left")
    C(p,f"C{r}",round(s['cn']/N['orders']),BOLD,None,RS,al="right");C(p,f"D{r}",round(s['cr']/R['orders']),BOLD,None,RS,al="right");C(p,f"E{r}",round((s['cn']+s['cr'])/tot_ord),BOLD,None,RS,al="right");C(p,f"F{r}","",bd=False); r+=1
r+=1
C(p,f"B{r}","Same margins & spend; only how media splits New/Repeat differs. Blended CM2 identical (~+2.9%).",ITAL,bd=False); r+=1
C(p,f"B{r}","Intent-based ≈ ceiling on new-media (Meta=acquisition); GA4 last-click ≈ floor (credits final click). Truth between.",ITAL,bd=False)

# ============================== DRIVERS ==============================
d=wb.create_sheet("Drivers & Sources"); d.sheet_view.showGridLines=False
for c,w in {"A":2,"B":44,"C":16,"D":46}.items(): d.column_dimensions[c].width=w
C(d,"B2","Drivers & sources (edit these, then recompute)",H1,navy,al="left",bd=False)
rows=[("KEY DRIVERS","",""),
 ("July TOTAL net revenue ₹",NETREV,"SCALE ANCHOR (estimate). COGS grand-total ÷ ~3 mo AND July media ÷ 43% both ≈ ₹8.3 Cr. Replace with DMR actual."),
 ("New-user share of net revenue",NEW_SHARE,"July-31 DMR sample (realized orders)."),
 ("AOV New ₹",AOV_NEW,"July-31 DMR sample."),
 ("AOV Repeat ₹",AOV_REP,"July-31 DMR sample."),
 ("Net Rev ÷ Gross Rev",NG,"COGS file grand total 247.9M/331.6M."),
 ("Net COGS % of Net Rev",COGS_R,"COGS file 111.14M/247.9M."),
 ("Fulfilment % of Net Rev",FUL_R,"COGS file 22.65M/247.9M."),
 ("CM1 % of Net Rev (derived)",CM1_R,"= 1 − COGS% − Fulfilment%."),
 ("MARKETING ACTUALS (July)","",""),
 ("Meta total ₹",META,"Windsor facebook. Prospecting ₹1.70 Cr + DPA/retarget ₹0.97 Cr."),
 ("Google total ₹",GOOGLE,"Windsor google_ads. PMax ₹0.82 Cr + Brand ₹0.05 Cr."),
 ("CRM (Moengage) ₹",CRM,"PLACEHOLDER — no rupee API (broadcast push≈free). Use invoice. →Repeat."),
 ("ATTRIBUTION SPLITS","",""),
 ("Meta new% — GA4 observed",META_NEW_GA4,"GA4 facebook instagram/paid: 45% new / 55% ret."),
 ("Google new% — GA4 observed",GOOG_NEW_GA4,"GA4 google/cpc: 57% new / 43% ret."),
]
rr=4
for lbl,val,src in rows:
    if val=="":
        C(d,f"B{rr}",lbl,WHITE,blue,al="left"); C(d,f"C{rr}","",fill=blue); C(d,f"D{rr}","",WHITE,blue,al="left")
    else:
        fmt=PCT if (isinstance(val,float) and val<1.5) else RS
        C(d,f"B{rr}",lbl,al="left"); C(d,f"C{rr}",val,BOLD,yell,fmt,al="right"); C(d,f"D{rr}",src,ITAL,al="wrap")
    rr+=1

# ============================== COGS REFERENCE ==============================
cr=wb.create_sheet("COGS Reference"); cr.sheet_view.showGridLines=False
for c,w in {"A":2,"B":26,"C":15,"D":14}.items(): cr.column_dimensions[c].width=w
C(cr,"B2","COGS/CM file — grand total (all months, all platforms)",H2,bd=False)
C(cr,"B3","Source: Web_Updated_cogs_CMFile.xlsx · Sheet2 pivot. Ratios drive the P&L.",ITAL,bd=False)
gt=[("Revenue",331557725),("Return Sales",43246094),("Net Sales",288311631),("Net Revenue",247899469),
    ("Net COGS",111138979),("Product Margin",136760490),("Fulfilment cost",22651338),("CM1",114109152),
    ("CM2",6509793),("Ad Spends",107599359),("QTY",207715)]
rr=5; C(cr,f"B{rr}","Metric",WHITE,blue,al="left");C(cr,f"C{rr}","Value ₹",WHITE,blue,al="center");C(cr,f"D{rr}","% Net Rev",WHITE,blue,al="center");rr+=1
for l,v in gt:
    C(cr,f"B{rr}",l,al="left"); C(cr,f"C{rr}",v,None,None,RS,al="right")
    C(cr,f"D{rr}",(v/247899469 if l!="QTY" else None),None,None,PCT if l!="QTY" else None,al="center"); rr+=1

# ============================== GA4 REFERENCE ==============================
g=wb.create_sheet("GA4 Attribution"); g.sheet_view.showGridLines=False
for c,w in {"A":2,"B":34,"C":14,"D":14,"E":16}.items(): g.column_dimensions[c].width=w
C(g,"B2","GA4 July — new vs returning (India acct 273672074)",H2,bd=False)
C(g,"B3","Session/last-click attribution. 'Returning'=visited before (NOT first-purchase). Windsor GA4.",ITAL,bd=False)
rr=5
C(g,f"B{rr}","Site-wide",WHITE,blue,al="left");C(g,f"C{rr}","Purch. rev ₹",WHITE,blue,al="center");C(g,f"D{rr}","Transactions",WHITE,blue,al="center");C(g,f"E{rr}","% of rev",WHITE,blue,al="center");rr+=1
for l,rev,tx in [("New",44232456,19035),("Returning",62674937,23288)]:
    C(g,f"B{rr}",l,al="left");C(g,f"C{rr}",rev,None,None,RS,al="right");C(g,f"D{rr}",tx,None,None,RS,al="right");C(g,f"E{rr}",rev/106907393,None,None,PCT,al="center");rr+=1
rr+=1
C(g,f"B{rr}","Paid channel mix (revenue)",WHITE,blue,al="left");C(g,f"C{rr}","New",WHITE,blue,al="center");C(g,f"D{rr}","Returning",WHITE,blue,al="center");C(g,f"E{rr}","",WHITE,blue);rr+=1
for l,nn,rrv in [("Meta — facebook instagram/paid",7977144,9636900),("Google — google/cpc",19815460,14986129)]:
    tot=nn+rrv
    C(g,f"B{rr}",l,al="left");C(g,f"C{rr}",nn/tot,None,grey,PCT,al="center");C(g,f"D{rr}",rrv/tot,None,grey,PCT,al="center");C(g,f"E{rr}","",bd=False);rr+=1
rr+=1
C(g,f"B{rr}","→ Meta spend splits 45/55; Google splits 57/43 → these feed the GA4 scenario in the P&L tab.",ITAL,bd=False)

# ============================== METHOD ==============================
m=wb.create_sheet("Method & Caveats"); m.sheet_view.showGridLines=False
m.column_dimensions["B"].width=112
notes=["METHOD & CAVEATS","",
 "Structural margins (COGS%, PM%, Fulfilment%, CM1/CM2%) — COGS file Sheet2 grand total (all-month blend).",
 "New-vs-Repeat mix (rev share, AOV) — july_DMR_unmasked.csv 'New User' flag, JULY-31 SAMPLE (816 realized orders).",
 "Meta & Google spend — Windsor.ai actuals, July. CRM — Moengage (no rupee API; placeholder).",
 "",
 "CONSTRAINT: 47MB DMR can't be fully aggregated (10MB download cap; text read truncates to July-31). So split ratios",
 "and the ₹8.3 Cr scale are directional. Replace with a full-month DMR pivot (New User × Category × Last Status →",
 "Σ Product Grand Total(FF), Σ discount, Σ Qty, distinct Order No.) to finalise exact numbers.",
 "",
 "ATTRIBUTION: New-vs-repeat SPEND split shown two ways. Intent-based = campaign type (prospecting→New, DPA→Repeat).",
 "GA4 last-click = each paid channel's observed new/returning purchase mix. They disagree on who's profitable because",
 "GA4 credits the final click (understating Meta's upper-funnel new-user role). Intent≈ceiling, GA4≈floor; truth between.",
 "",
 "COUPONS: Product Grand Total(FF) is already net of coupon → discounts sit ABOVE CM1 (in net revenue), NOT re-subtracted",
 "at CM2. Sample discount intensity: New ~19% / Repeat ~26% of realized revenue.",
 "",
 "VALIDATION: bottom-up blended CM2 (~+2.9%) ties to COGS file blended CM2 (2.6%).",]
rr=3
for n in notes:
    f=H1 if n=="METHOD & CAVEATS" else (BOLD if n[:9] in ("CONSTRAIN","ATTRIBUTI","COUPONS:","VALIDATIO") else ITAL)
    fl=navy if n=="METHOD & CAVEATS" else None; fo=Font(bold=True,size=14,color="FFFFFF") if n=="METHOD & CAVEATS" else f
    C(m,f"B{rr}",n,fo,fl,bd=False); rr+=1

path="/home/user/claude/out/DailyObjects_July2026_PnL_New_vs_Repeat.xlsx"
wb.save(path)
print("saved",path)
