"""
DailyObjects — Andromeda Creative Bank generator.
Replicates the structure/styling of the Loop reference workbook for every category.
Per category: Tab 1 = Persona Index (dominant-TG snapshot -> shortlisted personas),
              Tab 2 = Persona x Funnel Universe (each cell = a distinct creative).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = "Helvetica Neue"
DARK = "FF2B2B2B"      # header fill
WHITE = "FFFFFFFF"
BLACK = "FF000000"
# persona banding fills (from the reference, extended)
BANDS = ["FFE3ECF7", "FFFCEFE1", "FFE8F2E5", "FFF3E9F6", "FFDE9EE".rjust(8, "F"),
         "FFFBF6DD", "FFE6F4F7", "FFF0ECF9"]
BANDS = ["FFE3ECF7", "FFFCEFE1", "FFE8F2E5", "FFF3E9F6", "FFFDE9EE",
         "FFFBF6DD", "FFE6F4F7", "FFF0ECF9"]
SNAP_FILL = "FFF5F5F5"   # snapshot block fill
DOM_FILL = "FFEAF3EA"    # dominant tag
EXP_FILL = "FFFDEFE6"    # expansion tag

thin = Side(style="thin", color="FFDDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, text, ncols):
    c = ws.cell(1, 1, text)
    c.font = Font(name=FONT, size=13, bold=True, color=BLACK)
    c.alignment = Alignment(vertical="bottom", wrap_text=False)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 26


def header_row(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row, j, h)
        c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=DARK)
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 24


def band_of(idx):
    return BANDS[idx % len(BANDS)]


def data_cell(ws, r, c, val, fill=None, bold=False, wrap=True, color=BLACK, size=None):
    cell = ws.cell(r, c, val)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=wrap)
    cell.border = BORDER
    return cell


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
