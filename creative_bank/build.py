"""Assemble the DailyObjects Andromeda Creative Bank workbook."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from engine import (FONT, DARK, WHITE, BLACK, SNAP_FILL, DOM_FILL, EXP_FILL,
                    title, header_row, band_of, data_cell, set_widths, BORDER)

import cat_cases, cat_chargers, cat_organisers, cat_watchbands
import cat_backpacks, cat_laptopbags, cat_desks

CATS = [cat_cases.CATEGORY, cat_chargers.CATEGORY, cat_organisers.CATEGORY,
        cat_watchbands.CATEGORY, cat_backpacks.CATEGORY, cat_laptopbags.CATEGORY,
        cat_desks.CATEGORY]

FUNNEL_HEADERS = ["Persona", "Funnel step", "Awareness stage", "Mindset / trigger",
                  "Lead angle", "Positioning / message", "Structure", "Hook seed (first 3 sec)"]
PERSONA_HEADERS = ["Persona", "TG type", "Who they are", "Trigger / context", "Core need",
                   "Primary pain", "Hero angle", "Best SKU", "Where to reach (format / placement)"]

wb = openpyxl.Workbook()
wb.remove(wb.active)


# ---------- Framework key ----------
def build_framework():
    ws = wb.create_sheet("Framework Key")
    title(ws, "DAILYOBJECTS · ANDROMEDA CREATIVE BANK — HOW TO READ THIS UNIVERSE", 3)
    set_widths(ws, {"A": 34, "B": 72, "C": 96})
    header_row(ws, 3, ["Concept", "What it means", "Why it matters"])
    rows = [
        ("Two tabs per category",
         "Tab 1 (Personas) shows the dominant TG from that category's 3-year workbook and the 4-6 personas shortlisted from it. Tab 2 (Funnel Universe) maps every persona across the purchase journey — each cell is a distinct creative.",
         "Read the persona tab first to see WHO, then the funnel tab for WHAT to say to them at each step."),
        ("Persona = targeting",
         "Each category has 2-3 personas from the DOMINANT TG (who already buys) and 2-3 from the EXPANSION TG (the plan's growth bets — women, 35-44 professionals, premium Android, travellers, corporate).",
         "In the Andromeda era the creative finds the audience. A distinct persona-creative is read as a distinct signal — each persona unlocks a different slice of reach."),
        ("TG type — Dominant vs Expansion",
         "Dominant = the buyer the workbook's live Meta data already proves. Expansion = the under-indexed / new-TG segments the 3-year plan wants to win.",
         "Fund the dominant personas to defend and scale; fund the expansion personas to open the headroom the plan is built on."),
        ("Funnel step / Awareness",
         "TOFU->BOFU mapped to Unaware -> Problem-aware -> Solution-aware -> Product-aware -> Most-aware.",
         "The same persona needs a different message at each step. Don't show a BOFU offer to an Unaware viewer, or category education to someone ready to buy."),
        ("Angle ladder",
         "TOFU leads with Curiosity / Pain / Story. MOFU shifts to Desired outcome / Comparison / Authority. BOFU closes with Feature / Social proof / Offer.",
         "Diversifying the angle by stage is what stops your ads competing with each other and keeps CPMs down."),
        ("Structure",
         "PAS (Problem-Agitate-Solution), BAB (Before-After-Bridge), FAB (Feature-Advantage-Benefit), 4 U's (Useful / Urgent / Unique / Ultra-specific).",
         "Framing varies by stage so the algorithm treats each as a genuinely distinct concept, not a re-edit."),
        ("Hook seed",
         "The first 3 seconds — the single biggest lever for distinctness. Every cell has its own hook.",
         "6 personas x 5 stages = 30 distinct openers per category = 30 potential Entity IDs."),
        ("Dominant TG snapshot",
         "The block at the top of each Personas tab: gender / age / geography / performance, lifted from that category's live Meta campaign in the 3-year workbook.",
         "It grounds the personas in real buyer data — the personas are read OUT of this snapshot, not invented."),
        ("Use",
         "This is the universe, not the shoot list. Pick high-priority cells per quarter, brief them as distinct creatives, and slot into the bucket framework.",
         "Coverage > volume. Fill the empty persona x stage cells before re-shooting ones you already have."),
        ("Source",
         "Built from the DailyObjects category 3-year workbooks (FY27-29) and their live Meta TG data (via Windsor.ai). Reference structure: the Loop / Andromeda Creative Bank.",
         "Every category tab mirrors the Loop reference so the whole set reads as one system."),
    ]
    r = 4
    for a, b, c in rows:
        data_cell(ws, r, 1, a, bold=True)
        data_cell(ws, r, 2, b)
        data_cell(ws, r, 3, c)
        r += 1
    # category index
    r += 1
    data_cell(ws, r, 1, "CATEGORIES IN THIS WORKBOOK", bold=True, fill=DARK, color=WHITE)
    data_cell(ws, r, 2, "Dominant TG (gender · core age)", bold=True, fill=DARK, color=WHITE)
    data_cell(ws, r, 3, "Hero SKUs", bold=True, fill=DARK, color=WHITE)
    r += 1
    for cat in CATS:
        data_cell(ws, r, 1, cat["key"], bold=True, fill=band_of(CATS.index(cat)))
        data_cell(ws, r, 2, cat["snapshot"].get("Dominant TG — gender", "") + "  ·  " +
                  cat["snapshot"].get("Dominant TG — age", ""))
        data_cell(ws, r, 3, cat["snapshot"].get("Hero SKUs / price", ""))
        r += 1


def build_personas(cat, idx):
    ws = wb.create_sheet(f"{cat['tab']} · Personas")
    title(ws, f"{cat['title']} · PERSONA INDEX — Dominant TG -> Shortlisted Personas", len(PERSONA_HEADERS))
    set_widths(ws, {"A": 27, "B": 13, "C": 46, "D": 40, "E": 32, "F": 32, "G": 22, "H": 24, "I": 40})
    r = 2
    # snapshot block
    data_cell(ws, r, 1, "DOMINANT TG SNAPSHOT — from the 3-year workbook", bold=True, fill=DARK, color=WHITE)
    for j in range(2, len(PERSONA_HEADERS) + 1):
        c = ws.cell(r, j); c.fill = PatternFill("solid", fgColor=DARK); c.border = BORDER
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(PERSONA_HEADERS))
    r += 1
    for k, v in cat["snapshot"].items():
        data_cell(ws, r, 1, k, bold=True, fill=SNAP_FILL)
        data_cell(ws, r, 2, v, fill=SNAP_FILL)
        for j in range(3, len(PERSONA_HEADERS) + 1):
            c = ws.cell(r, j); c.fill = PatternFill("solid", fgColor=SNAP_FILL); c.border = BORDER
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(PERSONA_HEADERS))
        r += 1
    r += 1
    # persona table
    data_cell(ws, r, 1, "SHORTLISTED PERSONAS — 2-3 from Dominant TG + 2-3 from Expansion TG",
              bold=True, fill=DARK, color=WHITE)
    for j in range(2, len(PERSONA_HEADERS) + 1):
        c = ws.cell(r, j); c.fill = PatternFill("solid", fgColor=DARK); c.border = BORDER
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(PERSONA_HEADERS))
    r += 1
    header_row(ws, r, PERSONA_HEADERS)
    r += 1
    for pi, p in enumerate(cat["personas"]):
        band = band_of(pi)
        data_cell(ws, r, 1, p["name"], bold=True, fill=band)
        tg_fill = DOM_FILL if p["tg"] == "Dominant" else EXP_FILL
        data_cell(ws, r, 2, p["tg"], bold=True, fill=tg_fill)
        data_cell(ws, r, 3, p["who"], fill=band)
        data_cell(ws, r, 4, p["trigger"], fill=band)
        data_cell(ws, r, 5, p["need"], fill=band)
        data_cell(ws, r, 6, p["pain"], fill=band)
        data_cell(ws, r, 7, p["angle"], fill=band)
        data_cell(ws, r, 8, p["sku"], fill=band)
        data_cell(ws, r, 9, p["reach"], fill=band)
        r += 1
    ws.freeze_panes = ws.cell(row=[i for i in range(1, r)][0] + 0, column=1)  # placeholder
    ws.freeze_panes = None


def build_funnel(cat):
    ws = wb.create_sheet(f"{cat['tab']} · Funnel")
    title(ws, f"{cat['title']} · MAPPING UNIVERSE — Persona x Purchase Journey (each cell = a distinct creative)",
          len(FUNNEL_HEADERS))
    set_widths(ws, {"A": 27, "B": 11, "C": 16, "D": 40, "E": 24, "F": 82, "G": 10, "H": 56})
    header_row(ws, 2, FUNNEL_HEADERS)
    r = 3
    for pi, p in enumerate(cat["personas"]):
        band = band_of(pi)
        for row in p["funnel"]:
            step, stage, mindset, angle, positioning, structure, hook = row
            data_cell(ws, r, 1, p["name"], bold=True, fill=band)
            data_cell(ws, r, 2, step, fill=band, bold=True)
            data_cell(ws, r, 3, stage, fill=band)
            data_cell(ws, r, 4, mindset, fill=band)
            data_cell(ws, r, 5, angle, fill=band)
            data_cell(ws, r, 6, positioning, fill=band)
            data_cell(ws, r, 7, structure, fill=band)
            data_cell(ws, r, 8, hook, fill=band)
            r += 1
    ws.freeze_panes = "A3"


build_framework()
for i, cat in enumerate(CATS):
    build_personas(cat, i)
    build_funnel(cat)

out = "/home/user/claude/creative_bank/DailyObjects_CreativeBank_AllCategories.xlsx"
wb.save(out)
print("saved", out)
print("tabs:", wb.sheetnames)
