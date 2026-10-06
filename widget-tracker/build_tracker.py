"""Build the PDP/cart widget effectiveness tracker (BYS, CYP, YMAL).

Baseline actuals (30 Sep - 5 Oct 2026) were pulled from MoEngage on 6 Oct 2026.
Re-run to regenerate; daily actuals are then filled in by hand (blue cells).
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = "/home/user/claude/widget-tracker/DailyObjects_Widget_Effectiveness_Tracker.xlsx"
FONT = "Arial"
START, END = dt.date(2026, 9, 30), dt.date(2026, 12, 31)
DATES = [START + dt.timedelta(d) for d in range((END - START).days + 1)]
FIRST = 6                      # first data row on daily sheets
LAST = FIRST + len(DATES) - 1
BASE_ROWS = (FIRST, FIRST + 5)  # 30 Sep - 5 Oct baseline rows

# ---------- styles ----------
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=14, bold=True, color="1F3864")
f_sub = Font(name=FONT, size=10, italic=True, color="595959")
f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
fill_hdr = PatternFill("solid", fgColor="1F3864")
fill_band = {"site": "D9E1F2", "atc": "E2EFDA", "users": "FFF2CC", "orders": "FCE4D6",
             "rev": "EDE7F6", "guard": "F2F2F2", "status": "DDEBF7"}
fill_input = PatternFill("solid", fgColor="FFFF99")
fill_live = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap = Alignment(wrap_text=True, vertical="top")

INR = '"₹"#,##0;("₹"#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
DEC2 = '0.00;(0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = 'dd-mmm-yy'


def style_range(ws, rng, font=None, fill=None, fmt=None, border=True, align=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if border: c.border = box
            if align: c.alignment = align


wb = Workbook()

# =====================================================================
# Inputs
# =====================================================================
inp = wb.active
inp.title = "Inputs"
inp["A1"] = "Widget plan inputs & planned impact"; inp["A1"].font = f_title
inp["A2"] = ("Blue = editable input · Green = linked from another sheet · Black = formula. "
             "Yellow-filled cells are the plan levers to agree with the team.")
inp["A2"].font = f_sub
W = {"BYS": "C", "CYP": "D", "YMAL": "E"}
SHEET = {"BYS": "BYS Daily", "CYP": "CYP Daily", "YMAL": "YMAL Daily"}

hdr = ["Field", "Unit", "BYS", "CYP", "YMAL", "Source / note"]
for i, h in enumerate(hdr, 1):
    c = inp.cell(row=4, column=i, value=h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = box

rows = {}  # key -> row number
def irow(r, key, label, unit, vals, note="", kind="input", fmt=None):
    rows[key] = r
    inp.cell(row=r, column=1, value=label).font = f_base
    inp.cell(row=r, column=2, value=unit).font = f_sub
    for w, v in zip(["BYS", "CYP", "YMAL"], vals):
        c = inp[f"{W[w]}{r}"]
        c.value = v
        c.font = {"input": f_input, "link": f_link, "calc": f_base}[kind if not (isinstance(v, str) and v.startswith("=") is False and kind == "calc") else "input"]
        if kind == "input": c.fill = fill_input if key.startswith("lev_") else PatternFill()
        if fmt: c.number_format = fmt
        c.border = box
        c.alignment = Alignment(horizontal="right" if fmt else "left", vertical="center", wrap_text=True)
    inp.cell(row=r, column=6, value=note).font = f_sub
    inp.cell(row=r, column=6).alignment = wrap
    for col in (1, 2, 6): inp.cell(row=r, column=col).border = box

def section(r, title):
    inp.cell(row=r, column=1, value=title).font = f_bold
    for col in range(1, 7):
        inp.cell(row=r, column=col).fill = PatternFill("solid", fgColor="D9E1F2")

section(5, "1. Initiative definition & timeline")
irow(6, "name", "Widget", "", ["Build Your Stack", "CYP", "You May Also Like (new logic)"],
     "CYP full name / placement to be confirmed by product.")
irow(7, "place", "Placement", "", ["PDP", "TBC (assumed cart / pre-checkout)", "PDP"], "")
irow(8, "ident", "MoEngage identifier", "add_to_cart.referer",
     ["dp/quick-sell", "TBC - ask tech to tag a unique referer before go-live", "dp/cross-sell"],
     "BYS identified from today's post-launch spike (only referer that jumped after 1:20 PM). "
     "YMAL assumed to be the existing PDP cross-sell rail - confirm.")
irow(9, "golive", "Go-live date", "date",
     [dt.date(2026, 10, 6), dt.date(2026, 10, 8), dt.date(2026, 10, 15)],
     "BYS live 6-Oct ~1:20 PM. CYP Thu 8-Oct (Fri 9-Oct fallback - change this cell). "
     "YMAL logic mid-Oct. Every planned number on the daily sheets keys off this date.", fmt=DATE)
irow(10, "scope", "What changes at go-live", "",
     ["SKU recommendations added to BYS for a large share of PDPs (coverage expansion)",
      "New widget", "Recommendation logic upgrade on the existing rail"], "")

section(12, "2. Baseline (pre-launch) - links to 30-Sep to 5-Oct rows on each daily sheet")
def avg_link(w, col):
    s = SHEET[w]
    return f"=AVERAGE('{s}'!{col}{BASE_ROWS[0]}:{col}{BASE_ROWS[1]})"
def base_vals(col, zero_cyp=True):
    return [avg_link("BYS", col), 0 if zero_cyp else avg_link("CYP", col), avg_link("YMAL", col)]
irow(13, "b_events", "Widget ATC events / day", "events", base_vals("G"),
     "MoEngage behavior: add_to_cart, filter referer = identifier, daily. CYP = 0 (new widget).", "link", NUM1)
irow(14, "b_users", "Widget ATC users / day", "users", base_vals("J"),
     "MoEngage behavior, analysis type = users, same filter.", "link", NUM1)
irow(15, "b_value", "Widget ATC value / day", "₹", base_vals("Q"),
     "Sum of add_to_cart.value for widget ATCs (₹, item value added to cart).", "link", INR)
for w in ["BYS", "CYP", "YMAL"]:
    c = W[w]
    inp[f"{c}16"] = f"=IFERROR({c}15/{c}14,0)"
    inp[f"{c}17"] = (f"=IFERROR(SUM('{SHEET[w]}'!O{BASE_ROWS[0]}:O{BASE_ROWS[1]})/"
                     f"SUM('{SHEET[w]}'!J{BASE_ROWS[0]}:J{BASE_ROWS[1]}),0)") if w != "CYP" else 0
    inp[f"{c}18"] = f"={c}14*{c}17"
    inp[f"{c}19"] = f"={c}18*{c}16"
    inp[f"{c}20"] = f"=IFERROR({c}13/{c}14,0)"
irow(16, "b_vpu", "Value per widget ATC user", "₹", [inp["C16"].value, inp["D16"].value, inp["E16"].value], "", "calc", INR)
irow(17, "b_conv", "Widget ATC user → order (24h)", "%", [inp["C17"].value, inp["D17"].value, inp["E17"].value],
     "MoEngage funnel: add_to_cart (referer filter) → purchase, 24h window, unique users.", "link", PCT)
irow(18, "b_orders", "Widget-driven orders / day", "orders", [inp["C18"].value, inp["D18"].value, inp["E18"].value], "", "calc", NUM1)
irow(19, "b_attr", "Widget-attributed revenue / day", "₹", [inp["C19"].value, inp["D19"].value, inp["E19"].value],
     "Orders x value per ATC user (proxy for revenue of widget-added items in converted carts).", "calc", INR)
irow(20, "b_ipu", "ATC events per ATC user", "x", [inp["C20"].value, inp["D20"].value, inp["E20"].value], "", "calc", DEC2)

section(22, "3. Plan levers (steady state, after ramp-up)  - yellow = agree with team")
irow(23, "lev_mult", "ATC users multiplier vs baseline", "x", [1.6, None, 1.25],
     "BYS: coverage expansion; 6-Oct early read was +18% users / 2.3-3x events in first hours. "
     "YMAL: better relevance. CYP: n/a (new) - set target users directly in row 24.", "input", DEC2)
irow(24, "lev_users", "Target widget ATC users / day", "users", ["=C14*C23", 250, "=E14*E23"],
     "CYP assumption: ~2% of daily ATC users (~12k/day) add from CYP.", "input", NUM1)
irow(25, "lev_ipu", "Target ATC events per ATC user", "x", [1.6, 1.2, 1.4],
     "BYS should lift this (multi-item stacks). Baseline: BYS 1.35, YMAL 1.34.", "input", DEC2)
irow(26, "lev_conv", "Target ATC user → order (24h)", "%", [0.1125, 0.20, 0.055],
     "BYS held at baseline 11.3%. CYP is late-funnel so higher intent (assumed 20%). YMAL 4.6% → 5.5% with better relevance.", "input", PCT)
irow(27, "lev_vpu", "Target value per widget ATC user", "₹", [1800, 700, 1210],
     "BYS/YMAL ≈ baseline. CYP assumes accessory-led add-ons (~₹700).", "input", INR)
irow(28, "lev_incr", "Incrementality (share of uplift that is net new)", "%", [0.5, 0.6, 0.5],
     "Discount for cannibalisation of items users would have bought anyway. Replace with holdout read when available.", "input", PCT)
irow(29, "lev_ramp", "Ramp-up days to steady state", "days", [3, 5, 7],
     "Linear ramp from baseline to target over these days after go-live.", "input", NUM)
for w in W.values():  # make target users formula black for BYS/YMAL
    pass
inp["C24"].font = f_base; inp["E24"].font = f_base; inp["C24"].fill = PatternFill(); inp["E24"].fill = PatternFill()

section(31, "4. Planned impact per day at steady state")
for w in ["BYS", "CYP", "YMAL"]:
    c = W[w]
    inp[f"{c}32"] = f"={c}24*{c}25"
    inp[f"{c}33"] = f"={c}24*{c}26"
    inp[f"{c}34"] = f"={c}33*{c}27"
    inp[f"{c}35"] = f"={c}34-{c}19"
    inp[f"{c}36"] = f"={c}35*{c}28"
    inp[f"{c}37"] = f"={c}36*30"
    inp[f"{c}38"] = f"=IFERROR({c}36/$C$47,0)"
    inp[f"{c}39"] = f"=IFERROR({c}38/$C$49,0)"
    inp[f"{c}40"] = f"=IFERROR({c}36/$C$48,0)"
    inp[f"{c}41"] = f"=IFERROR({c}32/{c}13-1,0)" if w != "CYP" else '="new"'
plan_rows = [
    (32, "p_events", "Planned widget ATC events / day", "events", NUM),
    (33, "p_orders", "Planned widget-driven orders / day", "orders", NUM1),
    (34, "p_attr", "Planned widget-attributed revenue / day", "₹", INR),
    (35, "p_uplift", "Uplift in attributed revenue vs baseline / day", "₹", INR),
    (36, "p_net", "Planned NET incremental revenue / day", "₹", INR),
    (37, "p_net30", "Planned NET incremental revenue / 30 days", "₹", INR),
    (38, "p_aov", "Planned AOV uplift", "₹ per order", INR),
    (39, "p_aovpct", "Planned AOV uplift", "%", PCT),
    (40, "p_revpct", "Net incremental revenue as % of site revenue", "%", PCT),
    (41, "p_atcx", "Planned widget ATC events vs baseline", "%", PCT),
]
for r, key, label, unit, fmt in plan_rows:
    irow(r, key, label, unit, [inp[f"C{r}"].value, inp[f"D{r}"].value, inp[f"E{r}"].value],
         "AOV uplift assumes widget items attach to orders that would have happened anyway (no new orders)." if key == "p_aov" else "",
         "calc", fmt)
inp["F36"] = "Uplift x incrementality. This is the number to hold each widget to."

section(43, "5. Site baseline (30-Sep to 5-Oct, per day) - links to 'Site Daily'")
site_base = [
    (44, "PDP views / day", "events", "D", NUM), (45, "PDP viewers / day", "users", "E", NUM),
    (46, "All ATC users / day", "users", "F", NUM), (47, "Orders / day (purchase event)", "orders", "G", NUM),
    (48, "Revenue / day (purchase.value)", "₹", "H", INR),
]
for r, label, unit, col, fmt in site_base:
    inp.cell(row=r, column=1, value=label).font = f_base
    inp.cell(row=r, column=2, value=unit).font = f_sub
    c = inp[f"C{r}"]; c.value = f"=AVERAGE('Site Daily'!{col}{BASE_ROWS[0]}:{col}{BASE_ROWS[1]})"
    c.font = f_link; c.number_format = fmt
    for col_i in (1, 2, 3): inp.cell(row=r, column=col_i).border = box
inp["A49"] = "AOV (baseline)"; inp["B49"] = "₹"; inp["C49"] = "=IFERROR(C48/C47,0)"
inp["C49"].number_format = INR; inp["A49"].font = f_base; inp["B49"].font = f_sub
for col_i in (1, 2, 3): inp.cell(row=49, column=col_i).border = box
inp["F44"] = ("Baseline starts 30-Sep because add_to_cart / begin_checkout tracking changed on 29/30-Sep "
              "(ATC events ~3x overnight with flat PDP views). Pre-30-Sep data is not comparable.")
inp["F44"].font = Font(name=FONT, size=10, italic=True, color="C00000"); inp["F44"].alignment = wrap
inp.merge_cells("F44:F49")

section(51, "6. Portfolio total (all three widgets at steady state)")
tot = [(52, "Planned NET incremental revenue / day", "=SUM(C36:E36)", INR),
       (53, "Planned NET incremental revenue / 30 days", "=SUM(C37:E37)", INR),
       (54, "Planned AOV uplift (₹ per order)", "=SUM(C38:E38)", INR),
       (55, "Planned AOV uplift (%)", "=IFERROR(C54/C49,0)", PCT),
       (56, "Net incremental revenue as % of site revenue", "=IFERROR(C52/C48,0)", PCT)]
for r, label, f, fmt in tot:
    inp.cell(row=r, column=1, value=label).font = f_bold
    c = inp[f"C{r}"]; c.value = f; c.number_format = fmt; c.font = f_bold
    for col_i in (1, 2, 3): inp.cell(row=r, column=col_i).border = box

inp.column_dimensions["A"].width = 44; inp.column_dimensions["B"].width = 16
for c in "CDE": inp.column_dimensions[c].width = 24
inp.column_dimensions["F"].width = 70
for r in range(6, 11): inp.row_dimensions[r].height = 30
inp.freeze_panes = "C5"

# =====================================================================
# Site Daily
# =====================================================================
site = wb.create_sheet("Site Daily")
site["A1"] = "Site-level daily view - PDP funnel, revenue, AOV (planned vs actual)"; site["A1"].font = f_title
site["A2"] = ("Fill blue columns daily from MoEngage (day must be closed). Planned = baseline + sum of live widgets' "
              "planned net incremental revenue. Planned is flat (no weekday / festive seasonality) - read variance "
              "against the widget sheets, not just here. Diwali (8-Nov) will inflate actuals.")
site["A2"].font = f_sub
site_cols = [
    ("Date", "status", None), ("Day", "status", None), ("Widgets live", "status", NUM),
    ("PDP views", "site", NUM), ("PDP viewers", "site", NUM), ("All ATC users", "site", NUM),
    ("Orders", "site", NUM), ("Revenue (₹)", "site", INR), ("AOV (₹)", "site", INR),
    ("ATC users / PDP viewers", "site", PCT), ("Orders / PDP viewers", "site", PCT),
    ("Revenue / PDP viewer (₹)", "site", '"₹"#,##0.0'),
    ("Planned revenue (₹)", "rev", INR), ("Revenue var vs plan", "rev", PCT),
    ("Planned AOV (₹)", "rev", INR), ("AOV var vs plan", "rev", PCT),
    ("Widget net incr. rev - planned (₹)", "rev", INR), ("Widget net incr. rev - actual est. (₹)", "rev", INR),
    ("Main PDP ATC events (referer = dp)", "guard", NUM), ("Notes", "guard", None),
]
bands = [("A", "C", "Status", "status"), ("D", "L", "ACTUALS - site funnel (fill blue)", "site"),
         ("M", "R", "PLAN vs ACTUAL", "rev"), ("S", "T", "Guardrail", "guard")]
for a, b, t, k in bands:
    site.merge_cells(f"{a}4:{b}4"); c = site[f"{a}4"]; c.value = t; c.font = f_bold
    c.fill = PatternFill("solid", fgColor=fill_band[k]); c.alignment = center
for i, (h, band, fmt) in enumerate(site_cols, 1):
    c = site.cell(row=5, column=i, value=h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = box

site_actual = {
    "2026-09-30": (119492, 34505, 12602, 1379, 2644422.90, 18963),
    "2026-10-01": (120770, 34173, 12240, 1276, 2404727.29, 17733),
    "2026-10-02": (138639, 38042, 14006, 1508, 3166751.02, 20592),
    "2026-10-03": (131730, 34527, 11229, 1346, 2880679.67, 16637),
    "2026-10-04": (142710, 36847, 11940, 1417, 2913660.77, 17774),
    "2026-10-05": (123752, 34257, 11668, 1376, 2743496.28, 16430),
}
for i, d in enumerate(DATES):
    r = FIRST + i
    site[f"A{r}"] = d; site[f"A{r}"].number_format = DATE
    site[f"B{r}"] = f'=TEXT(A{r},"ddd")'
    site[f"C{r}"] = (f"=('BYS Daily'!C{r}=1)+('CYP Daily'!C{r}=1)+('YMAL Daily'!C{r}=1)")
    a = site_actual.get(d.isoformat())
    if a:
        for col, v in zip("DEFGHS", a): site[f"{col}{r}"] = v
    site[f"I{r}"] = f'=IF(G{r}="","",IFERROR(H{r}/G{r},""))'
    site[f"J{r}"] = f'=IF(E{r}="","",IFERROR(F{r}/E{r},""))'
    site[f"K{r}"] = f'=IF(E{r}="","",IFERROR(G{r}/E{r},""))'
    site[f"L{r}"] = f'=IF(E{r}="","",IFERROR(H{r}/E{r},""))'
    site[f"M{r}"] = f"=Inputs!$C$48+Q{r}"
    site[f"N{r}"] = f'=IF(H{r}="","",IFERROR(H{r}/M{r}-1,""))'
    site[f"O{r}"] = f"=IFERROR(M{r}/Inputs!$C$47,0)"
    site[f"P{r}"] = f'=IF(I{r}="","",IFERROR(I{r}/O{r}-1,""))'
    site[f"Q{r}"] = f"='BYS Daily'!U{r}+'CYP Daily'!U{r}+'YMAL Daily'!U{r}"
    site[f"R{r}"] = (f"=IF(AND('BYS Daily'!V{r}=\"\",'CYP Daily'!V{r}=\"\",'YMAL Daily'!V{r}=\"\"),\"\","
                     f"N('BYS Daily'!V{r})+N('CYP Daily'!V{r})+N('YMAL Daily'!V{r}))")
    for j, (h, band, fmt) in enumerate(site_cols, 1):
        c = site.cell(row=r, column=j); c.border = box
        c.font = f_input if L(j) in "DEFGHST" else (f_link if L(j) in "CQR" else f_base)
        if fmt: c.number_format = fmt
    site[f"A{r}"].font = f_base
site.freeze_panes = f"D{FIRST}"
widths = [11, 6, 8, 11, 11, 11, 9, 13, 10, 11, 11, 11, 14, 10, 11, 10, 14, 14, 13, 30]
for i, w in enumerate(widths, 1): site.column_dimensions[L(i)].width = w
site.row_dimensions[5].height = 45

# =====================================================================
# Widget daily sheets
# =====================================================================
wcols = [
    ("Date", "status", DATE), ("Day", "status", None), ("Live?", "status", NUM), ("Day # live", "status", NUM),
    ("Ramp factor", "status", '0%'),
    ("Planned", "atc", NUM), ("Actual", "atc", NUM), ("Var vs plan", "atc", PCT),
    ("Planned", "users", NUM), ("Actual", "users", NUM), ("Var vs plan", "users", PCT),
    ("Items per ATC user", "users", DEC2), ("ATC users per 1k PDP viewers", "users", NUM1),
    ("Planned", "orders", NUM1), ("Actual", "orders", NUM), ("ATC user→order %", "orders", PCT),
    ("Widget ATC value (₹) - actual", "rev", INR),
    ("Attributed revenue - planned (₹)", "rev", INR), ("Attributed revenue - actual (₹)", "rev", INR),
    ("Var vs plan", "rev", PCT),
    ("NET incremental rev - planned (₹)", "rev", INR), ("NET incremental rev - actual est. (₹)", "rev", INR),
    ("Widget ATC removals (if tracked)", "guard", NUM), ("Removal rate", "guard", PCT), ("Notes", "guard", None),
]
wbands = [("A", "E", "Status (auto from go-live date)", "status"), ("F", "H", "Widget ATC events", "atc"),
          ("I", "M", "Widget ATC unique users", "users"),
          ("N", "P", "Widget-driven orders (ATC user → purchase in 24h)", "orders"),
          ("Q", "V", "Revenue", "rev"), ("W", "Y", "Guardrail", "guard")]

widget_actual = {
    # date: (atc events, atc users, orders 24h, atc value)
    "BYS": {"2026-09-30": (469, 374, 51, 614233), "2026-10-01": (530, 379, 40, 707477),
            "2026-10-02": (703, 519, 49, 932400), "2026-10-03": (595, 427, 54, 744113),
            "2026-10-04": (579, 410, 43, 800023), "2026-10-05": (523, 406, 46, 719087)},
    "YMAL": {"2026-09-30": (1289, 970, 48, 1144216), "2026-10-01": (820, 595, 27, 764691),
             "2026-10-02": (1121, 819, 33, 1029688), "2026-10-03": (938, 700, 28, 843467),
             "2026-10-04": (855, 659, 31, 767251), "2026-10-05": (655, 495, 26, 587048)},
    "CYP": {},
}

for w in ["BYS", "CYP", "YMAL"]:
    ws = wb.create_sheet(SHEET[w])
    c = W[w]  # Inputs column
    ref = lambda key: f"Inputs!${c}${rows[key]}"
    ws["A1"] = f"{w} - daily planned vs actual"; ws["A1"].font = f_title
    ws["A2"] = (f'="Identifier: "&Inputs!{c}8&"   |   Go-live: "&TEXT(Inputs!{c}9,"dd-mmm-yy")&'
                f'"   |   Steady-state plan: "&TEXT(Inputs!{c}32,"#,##0")&" ATC events/day, ₹"&TEXT(Inputs!{c}36,"#,##0")&" net incremental revenue/day"')
    ws["A2"].font = f_sub
    ws["A3"] = ("Fill blue columns G, J, O, Q (and W if removals are tracked) daily from MoEngage once the day closes. "
                "Rows before go-live are the baseline. Green rows = widget live.")
    ws["A3"].font = f_sub
    for a, b, t, k in wbands:
        ws.merge_cells(f"{a}4:{b}4"); cc = ws[f"{a}4"]; cc.value = t; cc.font = f_bold
        cc.fill = PatternFill("solid", fgColor=fill_band[k]); cc.alignment = center
    for i, (h, band, fmt) in enumerate(wcols, 1):
        cc = ws.cell(row=5, column=i, value=h); cc.font = f_hdr; cc.fill = fill_hdr; cc.alignment = center; cc.border = box
    for i, d in enumerate(DATES):
        r = FIRST + i
        ws[f"A{r}"] = d
        ws[f"B{r}"] = f'=TEXT(A{r},"ddd")'
        ws[f"C{r}"] = f"=IF(A{r}>={ref('golive')},1,0)"
        ws[f"D{r}"] = f'=IF(C{r}=1,A{r}-{ref("golive")}+1,"")'
        ws[f"E{r}"] = f'=IF(C{r}=1,MIN(1,D{r}/{ref("lev_ramp")}),0)'
        ws[f"F{r}"] = f"={ref('b_events')}+({ref('p_events')}-{ref('b_events')})*E{r}"
        ws[f"H{r}"] = f'=IF(G{r}="","",IFERROR(G{r}/F{r}-1,""))'
        ws[f"I{r}"] = f"={ref('b_users')}+({ref('lev_users')}-{ref('b_users')})*E{r}"
        ws[f"K{r}"] = f'=IF(J{r}="","",IFERROR(J{r}/I{r}-1,""))'
        ws[f"L{r}"] = f'=IF(OR(G{r}="",J{r}=""),"",IFERROR(G{r}/J{r},""))'
        ws[f"M{r}"] = f"=IF(OR(J{r}=\"\",'Site Daily'!E{r}=\"\"),\"\",IFERROR(J{r}/'Site Daily'!E{r}*1000,\"\"))"
        ws[f"N{r}"] = f"=I{r}*({ref('b_conv')}+({ref('lev_conv')}-{ref('b_conv')})*E{r})"
        ws[f"P{r}"] = f'=IF(OR(O{r}="",J{r}=""),"",IFERROR(O{r}/J{r},""))'
        ws[f"R{r}"] = f"=N{r}*({ref('b_vpu')}+({ref('lev_vpu')}-{ref('b_vpu')})*E{r})"
        ws[f"S{r}"] = f'=IF(OR(O{r}="",Q{r}="",J{r}=""),"",IFERROR(O{r}*Q{r}/J{r},""))'
        ws[f"T{r}"] = f'=IF(S{r}="","",IFERROR(S{r}/R{r}-1,""))'
        ws[f"U{r}"] = f"=C{r}*(R{r}-{ref('b_attr')})*{ref('lev_incr')}"
        ws[f"V{r}"] = f'=IF(OR(C{r}=0,S{r}=""),"",(S{r}-{ref("b_attr")})*{ref("lev_incr")})'
        ws[f"X{r}"] = f'=IF(OR(W{r}="",G{r}=""),"",IFERROR(W{r}/G{r},""))'
        a = widget_actual[w].get(d.isoformat())
        if a:
            ws[f"G{r}"], ws[f"J{r}"], ws[f"O{r}"], ws[f"Q{r}"] = a
        for j, (h, band, fmt) in enumerate(wcols, 1):
            cc = ws.cell(row=r, column=j); cc.border = box
            col = L(j)
            cc.font = f_input if col in "GJOQWY" else f_base
            if fmt: cc.number_format = fmt
    ws.conditional_formatting.add(f"A{FIRST}:E{LAST}", FormulaRule(formula=[f"$C{FIRST}=1"], fill=fill_live))
    red = Font(name=FONT, size=10, color="C00000"); grn = Font(name=FONT, size=10, color="00803C")
    for col in "HKT":
        rng = f"{col}{FIRST}:{col}{LAST}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}{FIRST}),{col}{FIRST}<-0.1)'], font=red))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}{FIRST}),{col}{FIRST}>=0)'], font=grn))
    ws.freeze_panes = f"F{FIRST}"
    widths = [11, 6, 6, 7, 7, 10, 10, 9, 10, 10, 9, 9, 11, 10, 10, 10, 13, 14, 14, 9, 14, 14, 12, 9, 30]
    for i, wd in enumerate(widths, 1): ws.column_dimensions[L(i)].width = wd
    ws.row_dimensions[5].height = 48
    ws["G5"].comment = Comment("MoEngage behavior analysis: event add_to_cart, attribute filter referer = identifier "
                               "(Inputs row 8), granularity day, count = events.", "Tracker")
    ws["J5"].comment = Comment("Same query with analysis type = users (unique users).", "Tracker")
    ws["O5"].comment = Comment("MoEngage funnel: step1 add_to_cart (referer filter) -> step2 purchase, window 24h, "
                               "granularity day, unique users reaching step 2.", "Tracker")
    ws["Q5"].comment = Comment("MoEngage behavior, analysis type = aggregation, SUM of add_to_cart.value, referer filter.", "Tracker")
    ws["S5"].comment = Comment("Proxy: converting users x value added per ATC user. Replace with warehouse revenue of "
                               "order lines added via the widget when available.", "Tracker")

# =====================================================================
# Plan by Week
# =====================================================================
pw = wb.create_sheet("Plan by Week")
pw["A1"] = "Planned vs actual by week (Mon-Sun)"; pw["A1"].font = f_title
pw["A2"] = "All values are SUMIFS over the daily sheets. Actual columns only sum days that have actuals filled."
pw["A2"].font = f_sub
weeks = []
d = START - dt.timedelta(days=START.weekday())
while d <= END:
    weeks.append(d); d += dt.timedelta(7)
hdr2 = ["Week starting", "Week ending"]
blocks = []
for w in ["BYS", "CYP", "YMAL"]:
    blocks.append(w)
    hdr2 += [f"{w} ATC plan", f"{w} ATC actual", f"{w} net incr rev plan (₹)", f"{w} net incr rev actual (₹)"]
hdr2 += ["TOTAL net incr rev plan (₹)", "TOTAL net incr rev actual (₹)", "Site revenue actual (₹)", "Site AOV actual (₹)"]
for i, h in enumerate(hdr2, 1):
    cc = pw.cell(row=4, column=i, value=h); cc.font = f_hdr; cc.fill = fill_hdr; cc.alignment = center; cc.border = box
for k, wk in enumerate(weeks):
    r = 5 + k
    pw[f"A{r}"] = wk; pw[f"B{r}"] = f"=A{r}+6"
    col = 3
    for w in ["BYS", "CYP", "YMAL"]:
        s = f"'{SHEET[w]}'"
        rng = lambda cl: f"{s}!${cl}${FIRST}:${cl}${LAST}"
        crit = f'{rng("A")},">="&$A{r},{rng("A")},"<="&$B{r}'
        pw.cell(row=r, column=col, value=f"=SUMIFS({rng('F')},{crit})")
        pw.cell(row=r, column=col + 1, value=f'=SUMIFS({rng("G")},{crit})')
        pw.cell(row=r, column=col + 2, value=f"=SUMIFS({rng('U')},{crit})")
        pw.cell(row=r, column=col + 3, value=f'=SUMIFS({rng("V")},{crit})')
        col += 4
    pw.cell(row=r, column=col, value=f"=E{r}+I{r}+M{r}")
    pw.cell(row=r, column=col + 1, value=f"=F{r}+J{r}+N{r}")
    srng = lambda cl: f"'Site Daily'!${cl}${FIRST}:${cl}${LAST}"
    scrit = f'{srng("A")},">="&$A{r},{srng("A")},"<="&$B{r}'
    pw.cell(row=r, column=col + 2, value=f"=SUMIFS({srng('H')},{scrit})")
    pw.cell(row=r, column=col + 3, value=f"=IFERROR(SUMIFS({srng('H')},{scrit})/SUMIFS({srng('G')},{scrit}),0)")
    for j in range(1, len(hdr2) + 1):
        cc = pw.cell(row=r, column=j); cc.border = box; cc.font = f_base
        cc.number_format = DATE if j <= 2 else (NUM if (j - 3) % 4 in (0, 1) and j < 15 else INR)
r_tot = 5 + len(weeks)
pw[f"A{r_tot}"] = "TOTAL (30-Sep to 31-Dec)"; pw[f"A{r_tot}"].font = f_bold
for j in range(3, len(hdr2) + 1):
    cl = L(j)
    cc = pw.cell(row=r_tot, column=j, value=(f"=SUM({cl}5:{cl}{r_tot-1})" if j != len(hdr2) else
                                              f"=IFERROR(Q{r_tot}/SUM('Site Daily'!G{FIRST}:G{LAST}),0)"))
    cc.font = f_bold; cc.border = box; cc.number_format = pw.cell(row=5, column=j).number_format
pw.column_dimensions["A"].width = 12; pw.column_dimensions["B"].width = 12
for j in range(3, len(hdr2) + 1): pw.column_dimensions[L(j)].width = 14
pw.row_dimensions[4].height = 45
pw.freeze_panes = "C5"

# =====================================================================
# Summary (first tab)
# =====================================================================
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Widget effectiveness tracker - summary"; sm["A1"].font = f_title
sm["A2"] = ("Cumulative since each widget's go-live, counting only days with actuals filled, so planned and actual "
            "cover the same days. % to plan: green ≥100%, amber 80-99%, red <80%.")
sm["A2"].font = f_sub
shdr = ["Widget", "Placement", "Go-live", "Days live (with actuals)",
        "ATC events - plan", "ATC events - actual", "% to plan",
        "ATC users - plan", "ATC users - actual", "% to plan",
        "Widget orders - plan", "Widget orders - actual", "% to plan",
        "Attributed rev - plan (₹)", "Attributed rev - actual (₹)", "% to plan",
        "Net incr rev - plan (₹)", "Net incr rev - actual est. (₹)", "% to plan",
        "Steady-state net incr rev / day (₹)"]
for i, h in enumerate(shdr, 1):
    cc = sm.cell(row=4, column=i, value=h); cc.font = f_hdr; cc.fill = fill_hdr; cc.alignment = center; cc.border = box
for k, w in enumerate(["BYS", "CYP", "YMAL"]):
    r = 5 + k
    s = f"'{SHEET[w]}'"
    rg = lambda cl: f"{s}!${cl}${FIRST}:${cl}${LAST}"
    live_has = f'({rg("C")}=1)*({rg("G")}<>"")'
    sm[f"A{r}"] = f"=Inputs!{W[w]}6"
    sm[f"B{r}"] = f"=Inputs!{W[w]}7"
    sm[f"C{r}"] = f"=Inputs!{W[w]}9"
    sm[f"D{r}"] = f'=COUNTIFS({rg("C")},1,{rg("G")},"<>")'
    pairs = [("E", "F", "F", "G"), ("H", "I", "I", "J"), ("K", "L", "N", "O"), ("N", "O", "R", "S"), ("Q", "R", "U", "V")]
    for pc, ac, plan_col, act_col in pairs:
        sm[f"{pc}{r}"] = f'=SUMIFS({rg(plan_col)},{rg("C")},1,{rg("G")},"<>")'
        sm[f"{ac}{r}"] = f'=SUMIFS({rg(act_col)},{rg("C")},1,{rg("G")},"<>")'
        pct = L(ord(ac) - 64 + 1)
        sm[f"{pct}{r}"] = f'=IF({pc}{r}=0,"",{ac}{r}/{pc}{r})'
    sm[f"T{r}"] = f"=Inputs!{W[w]}36"
    for j in range(1, len(shdr) + 1):
        cc = sm.cell(row=r, column=j); cc.border = box
        cc.font = f_link if j <= 3 or j == 20 else f_base
        hd = shdr[j - 1]
        cc.number_format = (DATE if j == 3 else PCT if hd == "% to plan" else INR if "₹" in hd else NUM)
r = 8
sm[f"A{r}"] = "TOTAL"; sm[f"A{r}"].font = f_bold
for cl in "EFHIKLNOQRT":
    sm[f"{cl}{r}"] = f"=SUM({cl}5:{cl}7)"; sm[f"{cl}{r}"].font = f_bold
    sm[f"{cl}{r}"].number_format = sm[f"{cl}5"].number_format
for pc, ac, pct in [("E", "F", "G"), ("H", "I", "J"), ("K", "L", "M"), ("N", "O", "P"), ("Q", "R", "S")]:
    sm[f"{pct}{r}"] = f'=IF({pc}{r}=0,"",{ac}{r}/{pc}{r})'; sm[f"{pct}{r}"].number_format = PCT; sm[f"{pct}{r}"].font = f_bold
for j in range(1, 21): sm.cell(row=r, column=j).border = box
for cl in "GJMPS":
    rng = f"{cl}5:{cl}8"
    sm.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cl}5),{cl}5>=1)'], fill=PatternFill("solid", fgColor="C6EFCE")))
    sm.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cl}5),{cl}5>=0.8,{cl}5<1)'], fill=PatternFill("solid", fgColor="FFEB9C")))
    sm.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cl}5),{cl}5<0.8)'], fill=PatternFill("solid", fgColor="FFC7CE")))

# Site before/after block
sm["A11"] = "Site-level read (since first go-live vs baseline)"; sm["A11"].font = f_bold
sh = ["Metric", "Baseline / day (30-Sep to 5-Oct)", "Since 6-Oct / day (actual)", "Change", "Plan / day", "Var vs plan"]
for i, h in enumerate(sh, 1):
    cc = sm.cell(row=12, column=i, value=h); cc.font = f_hdr; cc.fill = fill_hdr; cc.alignment = center; cc.border = box
SD = "'Site Daily'"
post = lambda cl: (f'IFERROR(AVERAGEIFS({SD}!${cl}${FIRST}:${cl}${LAST},{SD}!$A${FIRST}:$A${LAST},">="&Inputs!$C$9,'
                   f'{SD}!${cl}${FIRST}:${cl}${LAST},"<>"),"")')
site_metrics = [
    ("PDP viewers", "=Inputs!C45", "=" + post("E"), None, NUM),
    ("Orders", "=Inputs!C47", "=" + post("G"), "=Inputs!C47", NUM),
    ("Revenue (₹)", "=Inputs!C48", "=" + post("H"), "=Inputs!C48+Inputs!C52", INR),
    ("AOV (₹)", "=Inputs!C49", '=IFERROR(C15/C14,"")', "=Inputs!C49+Inputs!C54", INR),
    ("Orders / PDP viewers", "=IFERROR(Inputs!C47/Inputs!C45,0)", '=IFERROR(C14/C13,"")', None, PCT),
    ("Revenue / PDP viewer (₹)", "=IFERROR(Inputs!C48/Inputs!C45,0)", '=IFERROR(C15/C13,"")', None, '"₹"#,##0.0'),
]
for k, (m, b, a, p, fmt) in enumerate(site_metrics):
    r = 13 + k
    sm[f"A{r}"] = m; sm[f"B{r}"] = b; sm[f"C{r}"] = a
    sm[f"D{r}"] = f'=IF(OR(C{r}="",B{r}=0),"",C{r}/B{r}-1)'
    if p: sm[f"E{r}"] = p; sm[f"F{r}"] = f'=IF(OR(C{r}="",E{r}=0),"",C{r}/E{r}-1)'
    for j in range(1, 7):
        cc = sm.cell(row=r, column=j); cc.border = box; cc.font = f_base
        cc.number_format = PCT if j in (4, 6) else fmt
sm["A20"] = ("Site AOV/revenue moves with traffic mix, offers and Diwali - use it as a sanity check only. "
             "Judge each widget on its own sheet (ATC users, orders, attributed revenue) and, ideally, a holdout.")
sm["A20"].font = f_sub
widths = [26, 20, 11, 10] + [12] * 15 + [14]
for i, wd in enumerate(widths, 1): sm.column_dimensions[L(i)].width = wd
sm.column_dimensions["B"].width = 22
sm.row_dimensions[4].height = 45; sm.row_dimensions[12].height = 32

# =====================================================================
# Metric dictionary & Read me
# =====================================================================
md = wb.create_sheet("Metric Dictionary")
md["A1"] = "Metric dictionary - what to track and why"; md["A1"].font = f_title
mh = ["Tier", "Metric", "Definition", "MoEngage query (event · filter · type)", "Why it matters / target"]
for i, h in enumerate(mh, 1):
    cc = md.cell(row=3, column=i, value=h); cc.font = f_hdr; cc.fill = fill_hdr; cc.alignment = center; cc.border = box
metrics = [
    ("1 · Primary", "Widget ATC events", "Add-to-carts fired from the widget", "add_to_cart · referer = identifier · behavior, events, daily", "Volume. Plan in Inputs row 32."),
    ("1 · Primary", "Widget ATC users", "Unique users who added at least 1 item from the widget", "same · behavior, users", "Reach. Separates 'more people' from 'same people adding more'."),
    ("1 · Primary", "Widget-driven orders", "Widget ATC users who purchased within 24h", "funnel add_to_cart(referer) → purchase · 24h · users", "Does widget intent convert? Plan in Inputs row 33."),
    ("1 · Primary", "Attributed revenue", "Converted widget users × value added per user (proxy)", "funnel + aggregation SUM(add_to_cart.value)", "₹ the widget touches. Swap for warehouse order-line revenue when available."),
    ("1 · Primary", "NET incremental revenue", "(Attributed rev − baseline attributed) × incrementality", "derived", "The number each widget is held to (Inputs row 36)."),
    ("1 · Primary", "AOV uplift", "Net incremental revenue ÷ site orders", "derived; site AOV = SUM(purchase.value)/purchases", "Primary business outcome for attach widgets."),
    ("2 · Diagnostic", "Items per ATC user", "Widget ATC events ÷ widget ATC users", "derived", "BYS should push this up (stacks)."),
    ("2 · Diagnostic", "ATC users per 1k PDP viewers", "Widget reach normalised for traffic", "view_item · users", "Removes traffic swings from the read."),
    ("2 · Diagnostic", "ATC user → order %", "Widget-driven orders ÷ widget ATC users", "funnel", "Falling = low-intent adds."),
    ("2 · Diagnostic", "Widget impressions & CTR", "Times the widget rendered / was clicked", "NOT TRACKED today - ask tech for a widget_view / widget_click event", "Needed to separate coverage from relevance."),
    ("2 · Diagnostic", "Orders containing a widget SKU / units per order", "Share of orders with ≥1 widget-added item", "Warehouse / OMS (needs line-level source tag)", "Best proof that adds turn into sold units."),
    ("3 · Guardrail", "Main-product ATC (referer = dp)", "PDP primary ATC button", "add_to_cart · referer = dp", "Should not fall - widget must not distract from the main product."),
    ("3 · Guardrail", "Widget ATC removal rate", "Removals of widget-added items ÷ widget ATCs", "remove_from_cart (if tracked) with source", "High = accidental / bulk adds (BYS risk)."),
    ("3 · Guardrail", "Site orders / PDP viewers, checkout completion", "Overall funnel health", "view_item, purchase", "Must not drop after go-live."),
    ("3 · Guardrail", "Returns / RTO on widget items", "Returned units from widget-added lines", "Warehouse", "Attach items with high returns erode the gain."),
]
for k, m in enumerate(metrics):
    for j, v in enumerate(m, 1):
        cc = md.cell(row=4 + k, column=j, value=v); cc.font = f_base; cc.alignment = wrap; cc.border = box
for i, wd in enumerate([14, 30, 45, 50, 50], 1): md.column_dimensions[L(i)].width = wd

rm = wb.create_sheet("Read Me", 1)
rm["A1"] = "How to use this tracker"; rm["A1"].font = f_title
notes = [
    ("Purpose", "Track the 3 recommendation widgets - Build Your Stack (BYS), CYP and You May Also Like (YMAL) - daily, planned vs actual, from each go-live date."),
    ("Tabs", "Summary (scorecard) · Inputs (go-live dates, baseline, plan levers, planned impact) · Site Daily · BYS / CYP / YMAL Daily · Plan by Week · Metric Dictionary."),
    ("Daily routine", "Next morning, for the closed day: fill blue cells on Site Daily (D-H, S) and on each live widget sheet (G, J, O, Q). Everything else calculates."),
    ("Colour legend", "Blue text = input · Green text = link to another sheet · Black = formula · Yellow fill = plan lever to agree · Green rows on widget sheets = widget live."),
    ("Changing a go-live", "Edit Inputs row 9. Planned numbers on every sheet shift automatically (with a linear ramp from Inputs row 29)."),
    ("Planned numbers", "Built from the 30-Sep to 5-Oct MoEngage baseline x plan levers (Inputs rows 23-29). The earlier per-widget sizing was not available in this session - overwrite the yellow levers with those numbers and the whole tracker updates."),
    ("Baseline caveat", "add_to_cart and begin_checkout tracking changed on 29/30-Sep (events ~3x with flat PDP views). The baseline uses only 30-Sep onwards. Do not compare against older data."),
    ("Open items", "1) CYP needs a unique referer before go-live, or it cannot be measured. 2) Confirm YMAL = dp/cross-sell. 3) Add widget impression/click events. 4) A 10-20% holdout per widget would replace the incrementality assumption with a measured number."),
    ("Data source", "MoEngage (workspace: DailyObjects). Queries for each column are in the header comments on the widget sheets and in the Metric Dictionary."),
]
for k, (a, b) in enumerate(notes):
    rm.cell(row=3 + k, column=1, value=a).font = f_bold
    cc = rm.cell(row=3 + k, column=2, value=b); cc.font = f_base; cc.alignment = wrap
    rm.row_dimensions[3 + k].height = 45
rm.column_dimensions["A"].width = 20; rm.column_dimensions["B"].width = 120

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
