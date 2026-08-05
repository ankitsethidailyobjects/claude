import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
d=json.load(open('/home/user/claude/delta.json'))
rows=d['rows']; tot=d['tot']; J=d['J']; A=d['A']
fn="/home/user/claude/DailyObjects_August_D2C_Plan.xlsx"
wb=openpyxl.load_workbook(fn)
if "6. Change vs Baseline" in wb.sheetnames: del wb["6. Change vs Baseline"]
s=wb.create_sheet("6. Change vs Baseline")
NAVY="1F2A44"; BLUE="2E5AAC"; LBLUE="DCE6F7"; GREY="F2F2F2"; WHITE="FFFFFF"
LGREEN="DDF0E1"; LRED="F7DEDE"; LAMBER="FBF0D5"
thin=Side(style='thin',color="C9C9C9"); bd=Border(*[thin]*4)
def st(c,b=False,sz=11,col="000000",fill=None,al='left',wrap=False,border=True,num=None):
    c.font=Font(bold=b,size=sz,color=col,name='Calibri')
    c.alignment=Alignment(horizontal=al,vertical='center',wrap_text=wrap)
    if fill:c.fill=PatternFill('solid',fgColor=fill)
    if border:c.border=bd
    if num:c.number_format=num
NUM='#,##0'; RS='"₹"#,##0'; PCT='0.0%'; RSCR='"₹"#,##0.00" cr"'; RSL='"₹"#,##0" L"'
def money_signed(v): return '"₹"+#,##0;"₹"-#,##0'
# title
s.merge_cells('A1:J1'); st(s.cell(1,1,"Change vs Baseline — June 2026 → August 2026 target"),True,16,WHITE,NAVY,border=False); s.row_dimensions[1].height=28
s.merge_cells('A2:J2'); st(s.cell(2,1,"June is the stated baseline month. Shows the shift needed: how many additional new/repeat orders & users, at what cost, from which categories."),False,10,"555555",border=False)

# Segment summary block
r=4; st(s.cell(r,1,"SEGMENT CHANGE (D2C, web+app)"),True,12,NAVY); r+=1
H=["Segment","June actual","August target","Δ orders","Δ %","Δ users (at plan CVR)","Note"]
for i,h in enumerate(H,1): st(s.cell(r,i,h),True,10,WHITE,BLUE,'left' if i in(1,7) else 'center',wrap=True)
r+=1
seg=[("New orders",J['newO'],A['newO'],A['newO']-J['newO'],(A['newO']/J['newO']-1),A['newU']-J['newU'],"Sessions FALL: new CVR normalises 0.46%→0.63%"),
     ("Repeat orders",J['repO'],A['repO'],A['repO']-J['repO'],(A['repO']/J['repO']-1),A['repU']-J['repU'],"The real volume engine for August"),
     ("TOTAL orders",J['newO']+J['repO'],A['newO']+A['repO'],(A['newO']+A['repO'])-(J['newO']+J['repO']),((A['newO']+A['repO'])/(J['newO']+J['repO'])-1),(A['newU']+A['repU'])-(J['newU']+J['repU']),"")]
for name,jv,av,dv,dp,du,note in seg:
    isT = name.startswith("TOTAL")
    fill=NAVY if isT else (GREY if r%2 else WHITE); fc=WHITE if isT else "000000"
    st(s.cell(r,1,name),True,10,fc,fill,'left')
    for i,(v,f) in enumerate([(jv,NUM),(av,NUM),(dv,'+#,##0;-#,##0'),(dp,'+0.0%;-0.0%'),(du,'+#,##0;-#,##0')],2):
        c=s.cell(r,i,v); st(c,isT,10,fc,fill,'center'); c.number_format=f
    st(s.cell(r,7,note),False,9,fc if isT else "555555",fill,'left',wrap=True); r+=1
r+=1
st(s.cell(r,1,"Revenue: ₹9.09 cr → ₹10.50 cr  (+₹1.41 cr, +15.5%). ~₹1.26 cr of the +₹1.41 cr comes from repeat. Incremental Meta spend to fund the shift ≈ ₹19 L."),False,10,NAVY,LBLUE,'left',wrap=True)
s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7); s.row_dimensions[r].height=30

# Category change table
r+=2; st(s.cell(r,1,"CATEGORY CHANGE — where the additional orders & users come from"),True,12,NAVY); r+=1
H=["Category","LTV","New Jun→Aug","Δ New ord","Δ New users","Repeat Jun→Aug","Δ Rep ord","Δ Rep users","Δ Spend New","Δ Spend Rep"]
for i,h in enumerate(H,1): st(s.cell(r,i,h),True,9,WHITE,BLUE,'left' if i<=2 else 'center',wrap=True)
r+=1
for x in rows:
    fill=LGREEN if x['ltv']=='HIGH' else (LAMBER if x['ltv']=='LOW' else (GREY if r%2 else WHITE))
    st(s.cell(r,1,x['k']),True,9,al='left',fill=fill)
    st(s.cell(r,2,x['ltv']),False,9,al='center',fill=fill)
    st(s.cell(r,3,f"{x['jN']:,.0f} → {x['aN']:,.0f}"),False,9,al='center',fill=fill)
    c=s.cell(r,4,round(x['dN'])); st(c,False,9,al='center',fill=fill); c.number_format='+#,##0;-#,##0'
    c=s.cell(r,5,round(x['dNu'])); st(c,False,9,al='center',fill=fill); c.number_format='+#,##0;-#,##0'
    st(s.cell(r,6,f"{x['jR']:,.0f} → {x['aR']:,.0f}"),False,9,al='center',fill=fill)
    c=s.cell(r,7,round(x['dR'])); st(c,False,9,al='center',fill=fill); c.number_format='+#,##0;-#,##0'
    c=s.cell(r,8,round(x['dRu'])); st(c,False,9,al='center',fill=fill); c.number_format='+#,##0;-#,##0'
    c=s.cell(r,9,round(x['dNs'])); st(c,False,9,al='center',fill=fill); c.number_format='"₹"+#,##0;"₹"-#,##0'
    c=s.cell(r,10,round(x['dRs'])); st(c,False,9,al='center',fill=fill); c.number_format='"₹"+#,##0;"₹"-#,##0'
    r+=1
# total
vals=["TOTAL","",f"{tot['jN']:,.0f} → {tot['aN']:,.0f}",tot['dN'],tot['dNu'],f"{tot['jR']:,.0f} → {tot['aR']:,.0f}",tot['dR'],tot['dRu'],tot['dNs'],tot['dRs']]
for i,v in enumerate(vals,1):
    c=s.cell(r,i,round(v) if isinstance(v,float) else v)
    st(c,True,9,WHITE,NAVY,'left' if i<=2 else 'center')
    if i in(4,5,7,8): c.number_format='+#,##0;-#,##0'
    if i in(9,10): c.number_format='"₹"+#,##0;"₹"-#,##0'
r+=2
notes=["Read: grow HIGH-LTV (Stack, Power Bank, Charging Station) on both new & repeat; hold MED; cut Watch Bands (low LTV/ROAS).",
 "Δ New users / Δ Rep users = additional sessions to acquire for the incremental orders, at plan CVR (new 0.63%, repeat 1.53%).",
 "Δ Spend = incremental Meta media: new = Δord × category CAC; repeat = Δord × 35% paid × ₹865. Net new +₹4.9L, repeat +₹14.1L.",
 "Repeat additions run via CRM + app push + retargeting — most of the 1.43M repeat sessions are owned-channel, not paid."]
for n in notes:
    s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
    st(s.cell(r,1,n),False,9,"333333",border=False,al='left',wrap=True); s.row_dimensions[r].height=16; r+=1
for i,wd in enumerate([23,7,16,11,12,16,11,12,12,12],1): s.column_dimensions[get_column_letter(i)].width=wd
wb.save(fn)
print("added tab. sheets:",wb.sheetnames)
