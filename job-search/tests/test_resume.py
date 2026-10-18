import io

import pytest
from docx import Document

from jobsearch.resume import ResumeError, extract_text

BODY = "Experienced category manager with D2C and marketplace P&L ownership. " * 6


def test_txt():
    assert "category manager" in extract_text("cv.txt", BODY.encode())


def test_docx_includes_tables():
    doc = Document()
    doc.add_paragraph(BODY)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "SQL"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("cv.docx", buf.getvalue())
    assert "category manager" in text and "Skills | SQL" in text


def test_rejects_unknown_format():
    with pytest.raises(ResumeError):
        extract_text("cv.png", b"...")


def test_rejects_too_little_text():
    with pytest.raises(ResumeError):
        extract_text("cv.txt", b"hi")
