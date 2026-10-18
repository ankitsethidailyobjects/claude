"""Extract plain text from an uploaded resume."""

import io
import re
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class ResumeError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise ResumeError(f"Unsupported resume format: {suffix or 'unknown'} (use PDF, DOCX or TXT)")

    text = _clean(text)
    if len(text) < 200:
        raise ResumeError(
            "Could not read enough text from the resume. If it is a scanned PDF, "
            "upload a DOCX or a text-based PDF instead."
        )
    return text


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()
