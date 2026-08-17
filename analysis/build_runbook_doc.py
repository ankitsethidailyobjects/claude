#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor, Inches
SEG="analysis/seg"; OUT=f"{SEG}/DailyObjects_Cohort_Refresh_Runbook.docx"
BLUE=RGBColor(0x1e,0x40,0xaf); INK=RGBColor(0x0f,0x1b,0x2d); GREY=RGBColor(0x5b,0x6b,0x7f)
d=Document(); st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5); st.font.color.rgb=INK
def H(t,s=15,before=12,after=6,color=BLUE):
    p=d.add_paragraph();p.space_before=Pt(before);p.space_after=Pt(after)
    r=p.add_run(t);r.bold=True;r.font.size=Pt(s);r.font.color.rgb=color;return p
def para(t,size=10.5,color=INK,after=6,bold=False):
    p=d.add_paragraph();p.space_after=Pt(after);r=p.add_run(t);r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=color;return p
def bullet(t,lead=None):
    p=d.add_paragraph(style="List Bullet");p.space_after=Pt(3)
    if lead:p.add_run(lead+" ").bold=True
    p.add_run(t);return p
def numbered(t,lead=None):
    p=d.add_paragraph(style="List Number");p.space_after=Pt(3)
    if lead:p.add_run(lead+" ").bold=True
    p.add_run(t);return p
def mono(t):
    p=d.add_paragraph();p.space_after=Pt(4);r=p.add_run(t);r.font.name="Consolas";r.font.size=Pt(9);r.font.color.rgb=RGBColor(0x33,0x41,0x55);return p
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.style="Light Grid Accent 1"
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i];c.text="";r=c.paragraphs[0].add_run(h);r.bold=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text="";rr=cs[i].paragraphs[0].add_run(str(v));rr.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    d.add_paragraph().space_after=Pt(2);return t

p=d.add_paragraph();r=p.add_run("DailyObjects — Cohort Refresh & MoEngage Runbook");r.bold=True;r.font.size=Pt(21);r.font.color.rgb=INK
s=d.add_paragraph();rr=s.add_run("Operating rules for refreshing the customer cohort system and syncing to MoEngage");rr.italic=True;rr.font.size=Pt(11);rr.font.color.rgb=GREY
d.add_paragraph()

H("1. Principle",16)
para("The warehouse is the source of truth and does all computation; MoEngage is the activation layer that receives the finished decisions. Identity layers (lifecycle & behavioural cohort) change slowly; state layers (activity, recency) change fast. Refresh each on its own cadence — never re-cluster the whole base daily.")

H("2. Refresh cadence matrix",16)
table(["Signal / layer","Refresh","Why"],
 [["Lifetime order count","Daily (event-driven)","Changes only on a new order"],
  ["Lifecycle cohort (O1 / O2 / 3+)","Daily","Deterministic; must flip the day a customer crosses to 3+"],
  ["Recency / days-since-last","Daily","Increments daily; drives activity & triggers"],
  ["Activity state (Active/Lapsed/Dormant)","Daily","Fast-moving"],
  ["Category & sub-category affinity","Daily on order (recompute weekly)","Cheap; drives creative selection"],
  ["Discount class / value tier","Weekly","Slow-moving ratios; daily churn is noise"],
  ["Behavioural cohort (re-score to centroids)","Monthly","Matches CRM planning cycle & cluster stability; hysteresis-guarded"],
  ["Behavioural model + event windows (re-train)","Quarterly (or on drift)","Re-fitting centroids & recalibrating price-led event depths is identity-changing"]],
 widths=[2.6,1.9,2.1])

H("3. The pipeline (run order)",16)
para("All scripts live in analysis/. A monthly refresh with a new DMR month runs top to bottom:",color=GREY,size=10)
numbered("Land the new month's DMR file(s) into data/raw/ (download from Drive).","Ingest —")
numbered("Extract the new file: extract_features.py (RFM/discount/category) and extract_subcat.py (sub-category).","Extract —")
numbered("step1_resolve.py — re-resolve identity across the full base (email+phone+UUID, hub-suppressed).","Resolve —")
numbered("enrich_discounts.py — recompute weekly-ASP event windows and per-order price-led + coupon discount.","Discounts —")
numbered("step2_features.py 2026-08-31 prod  +  step6_3cohorts.py — rebuild features and re-score the 3 cohorts (monthly).","Cohorts —")
numbered("build_repository.py — regenerate user_repository.csv (the MoEngage-ready master).","Repository —")
numbered("Push user_repository.csv to MoEngage (Section 5).","Sync —")

H("4. How an analyst refreshes it via Claude",16)
para("No code required. Open a Claude Code session on this repository and give the instruction; Claude runs the pipeline and reports back. Example prompts:")
mono('“Refresh the cohort model with the latest DMR through August 2026: add the new DMR file from Drive, run the full pipeline (extract → resolve → enrich discounts → re-score 3 cohorts → rebuild user_repository.csv), and show me the updated cohort sizes and retention.”')
mono('“Re-train the behavioural model this quarter, recalibrate the price-led event windows, and give me a V1→V2 cohort migration crosswalk before we cut over.”')
mono('“Export the MoEngage upload file for the current cohorts and push it as user attributes.”')
para("Claude will: verify the new data, execute the scripts in order (Section 3), sanity-check cohort sizes vs last run, regenerate the repository, and (on approval) sync to MoEngage. Guardrail: Claude flags if cohort sizes move more than ±20% month-on-month before pushing.",size=10,color=GREY)

H("5. MoEngage upload",16)
para("Push only the finished decisions as USER ATTRIBUTES (keyed on email / phone), plus a few EVENTS for journey triggers. Do not push raw ML features.")
table(["MoEngage field","Type","Source col","Refresh"],
 [["lifecycle_cohort","attr","cohort / order_cohort","daily"],
  ["behavioural_cohort","attr","cohort (Power/Deal/Light)","monthly"],
  ["activity_state","attr","activity_state","daily"],
  ["value_tier","attr","value_tier","weekly"],
  ["cat_affinity + tech_score","attr","cat_affinity, tech_score","weekly"],
  ["primary_subcategory / secondary","attr","primary_subcategory,…","weekly"],
  ["days_since_last, lifetime_orders, aov","attr","(same)","daily/weekly"],
  ["cohort_changed / became_dormant / reactivated / reorder_due","event","derived deltas","near-real-time / daily"]],
 widths=[2.4,1.0,2.4,1.2])
para("Method: bulk user-attribute import via CSV to MoEngage (S3/SFTP connector or the User Import API); events via the MoEngage events API/stream. Full-file attribute refresh daily for fast fields; monthly for behavioural_cohort. Segments are built INSIDE MoEngage by combining these attributes (e.g. Deal Seekers + Dormant + Tech affinity).",size=10)

H("6. Versioning",16)
bullet("every cohort assignment carries model_version, cohort_version and assignment_date.","Stamp —")
bullet("keep an SCD-2 history table so cohort migration over time is queryable; never overwrite old assignments.","History —")
bullet("run V2 in shadow for one cycle, publish a V1→V2 crosswalk, then cut MoEngage over on a communicated date.","New model —")
para("")
f=d.add_paragraph();fr=f.add_run("Pipeline scripts: extract_features.py · extract_subcat.py · step1_resolve.py · enrich_discounts.py · step2_features.py · step6_3cohorts.py · build_repository.py");fr.italic=True;fr.font.size=Pt(8);fr.font.color.rgb=GREY
d.save(OUT); print("saved",OUT)
