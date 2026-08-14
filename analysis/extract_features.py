#!/usr/bin/env python3
"""
Rich extractor for behavioural features. Emits ONE row per (order_no, category)
with order-level scalars carried, so downstream we can build both order-level
and user-level category features.

Output columns:
  order_no, order_date, paid, delivered, email, phone, uuid,
  payment, source, category, cat_gmv, cat_qty, cat_disc, order_disc, coupon
"""
import sys, os, csv, re, datetime as dt
from collections import defaultdict

NEVER_PAID={"paymentpending","paymentfailed"}
FULFILLED={"delivered","shipped","reversepickupdone","reversepickupinitiated","rto","invoiced","cod","codverified","paymentreceived"}
EMAIL_PH={"","na","n/a","none","null","-","useremail","user email","guest","test","nan"}
UUID_RE=re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$")

def _norm(h): return re.sub(r"[^a-z0-9]","",str(h).strip().lower()) if h is not None else ""
def nemail(v):
    if v is None:return ""
    s=str(v).strip().lower()
    return "" if (s in EMAIL_PH or "@" not in s or "." not in s.split("@")[-1]) else s
def nphone(v):
    if v is None:return ""
    s=str(v).strip()
    if s.endswith(".0"):s=s[:-2]
    d=re.sub(r"\D","",s)
    if len(d)>=12 and d.startswith("91"):d=d[-10:]
    elif len(d)==11 and d.startswith("0"):d=d[-10:]
    elif len(d)>10:d=d[-10:]
    if len(d)!=10 or d[0] not in "6789" or len(set(d))<=2:return ""
    return d
def nuuid(v):
    if v is None:return ""
    s=str(v).strip().lower()
    return s if UUID_RE.match(s) else ""
def pdate(v):
    if v is None:return None
    if isinstance(v,dt.datetime):return v.date()
    if isinstance(v,dt.date):return v
    s=str(v).strip().split(" ")[0]
    if not s:return None
    for f in ("%m/%d/%Y","%Y-%m-%d","%d/%m/%Y","%m/%d/%y","%d-%m-%Y"):
        try:return dt.datetime.strptime(s,f).date()
        except ValueError:pass
    return None
def num(v):
    if v is None:return 0.0
    try:return float(str(v).replace(",","").strip() or 0)
    except:return 0.0

def has_cols(header,strict=True):
    s={_norm(h) for h in header}
    ok="orderno" in s and ("useremail" in s or "email" in s)
    if not strict:return ok
    return ok and ("onlydate" in s or "orderdate" in s) and ("itemstatus" in s or "laststatus" in s)

def rows(path):
    ext=os.path.splitext(path)[1].lower()
    if ext==".csv":
        with open(path,newline="",encoding="utf-8-sig",errors="replace") as f:
            for r in csv.reader(f):yield r
    else:
        import openpyxl
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        firsts={}
        for sn in wb.sheetnames:
            try:firsts[sn]=next(wb[sn].iter_rows(values_only=True))
            except StopIteration:firsts[sn]=[]
        ch=(next((sn for sn in wb.sheetnames if has_cols(firsts[sn],True)),None)
            or next((sn for sn in wb.sheetnames if has_cols(firsts[sn],False)),None) or wb.sheetnames[0])
        sys.stderr.write(f"[sheet]{ch}\n")
        for r in wb[ch].iter_rows(values_only=True):yield list(r)

def idx(header,*names):
    hl=[_norm(h) for h in header]
    for n in names:
        if _norm(n) in hl:return hl.index(_norm(n))
    return None

def main():
    path,out=sys.argv[1],sys.argv[2]
    it=rows(path);header=next(it)
    i_ord=idx(header,"Order No.");i_em=idx(header,"User Email","Email")
    i_od=idx(header,"Only Date","Order Date");i_od2=idx(header,"Order Date")
    i_st=idx(header,"Item Status","Last Status")
    i_ph=idx(header,"Login Phone Number","Mobile","Phone");i_ph2=idx(header,"Mobile","Phone")
    i_uid=idx(header,"Username")
    i_pay=idx(header,"Payment Method");i_src=idx(header,"Source")
    i_cat=idx(header,"Category")
    i_gmv=idx(header,"Product Grand Total(FF)","Product Grand Total","Product final Price")
    i_qty=idx(header,"Qty Ordered")
    i_pdisc=idx(header,"Product Discount");i_mdisc=idx(header,"Membership Discount")
    i_odisc=idx(header,"Order Discount");i_coup=idx(header,"Coupon Code")
    assert None not in (i_ord,i_em,i_od,i_st),f"missing {header[:5]}"
    need=[x for x in (i_ord,i_em,i_od,i_od2,i_st,i_ph,i_ph2,i_uid,i_pay,i_src,i_cat,i_gmv,i_qty,i_pdisc,i_mdisc,i_odisc,i_coup) if x is not None]
    maxc=max(need)

    oscalar={}   # order -> [date,paid,deliv,email,phone,uuid,pay,src,order_disc,coupon]
    ocat=defaultdict(lambda:[0.0,0.0,0.0])  # (order,cat)->[gmv,qty,disc]
    for r in it:
        if not r or len(r)<=maxc:continue
        o=r[i_ord]
        if o is None or str(o).strip()=="":continue
        o=str(o).strip()
        d=pdate(r[i_od]) or (pdate(r[i_od2]) if i_od2 is not None else None)
        st=str(r[i_st]).strip().lower() if r[i_st] is not None else ""
        paid=0 if st in NEVER_PAID else 1
        deliv=1 if st in FULFILLED else 0
        em=nemail(r[i_em]);ph=nphone(r[i_ph]) or (nphone(r[i_ph2]) if i_ph2 is not None else "")
        uid=nuuid(r[i_uid]) if i_uid is not None else ""
        pay=str(r[i_pay]).strip().lower() if i_pay is not None and r[i_pay] is not None else ""
        src=str(r[i_src]).strip().lower() if i_src is not None and r[i_src] is not None else ""
        odisc=num(r[i_odisc]) if i_odisc is not None else 0.0
        coup=1 if (i_coup is not None and r[i_coup] not in (None,"")) else 0
        cat=str(r[i_cat]).strip() if i_cat is not None and r[i_cat] is not None and str(r[i_cat]).strip() else "Unknown"
        gmv=num(r[i_gmv]) if i_gmv is not None else 0.0
        qty=num(r[i_qty]) if i_qty is not None else 1.0
        ldisc=(num(r[i_pdisc]) if i_pdisc is not None else 0.0)+(num(r[i_mdisc]) if i_mdisc is not None else 0.0)

        s=oscalar.get(o)
        if s is None:
            oscalar[o]=[d,paid,deliv,em,ph,uid,pay,src,odisc,coup]
        else:
            if d and (s[0] is None or d<s[0]):s[0]=d
            s[1]|=paid;s[2]|=deliv
            if not s[3] and em:s[3]=em
            if not s[4] and ph:s[4]=ph
            if not s[5] and uid:s[5]=uid
            if not s[6] and pay:s[6]=pay
            if not s[7] and src:s[7]=src
            s[8]=max(s[8],odisc);s[9]|=coup
        c=ocat[(o,cat)];c[0]+=gmv;c[1]+=qty;c[2]+=ldisc

    with open(out,"w",newline="") as f:
        w=csv.writer(f)
        w.writerow(["order_no","order_date","paid","delivered","email","phone","uuid",
                    "payment","source","category","cat_gmv","cat_qty","cat_disc","order_disc","coupon"])
        for (o,cat),(gmv,qty,disc) in ocat.items():
            s=oscalar[o]
            w.writerow([o, s[0].isoformat() if s[0] else "", s[1], s[2], s[3], s[4], s[5],
                        s[6], s[7], cat, round(gmv,2), int(qty), round(disc,2), round(s[8],2), s[9]])
    sys.stderr.write(f"orders={len(oscalar)} ordercat_rows={len(ocat)}\n")
    print(f"{os.path.basename(path)}: orders={len(oscalar)} rows={len(ocat)}")

if __name__=="__main__":main()
