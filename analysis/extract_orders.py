#!/usr/bin/env python3
"""
Extract unique ORDERS from a DMR file, capturing all user identifiers so that
customers can later be resolved across email masking / year boundaries.

DMR is order-LINE level (one row per SKU). We aggregate each Order No. to one
record and keep: email, phone (mobile), uuid (account Username), date, status.

Output CSV columns:
  order_no, order_date(YYYY-MM-DD), year, paid, delivered, email, phone, uuid
    paid      = 1 if any line status != paymentPending/paymentFailed
    delivered = 1 if any line is a fulfilled status
"""
import sys, os, csv, json, re, datetime as dt
from collections import Counter

NEVER_PAID = {"paymentpending", "paymentfailed"}
FULFILLED  = {"delivered", "shipped", "reversepickupdone", "reversepickupinitiated",
              "rto", "invoiced", "cod", "codverified", "paymentreceived"}
EMAIL_PLACEHOLDERS = {"", "na", "n/a", "none", "null", "-", "useremail", "user email",
                      "guest", "test", "nan"}
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$")

def _norm(h):
    if h is None: return ""
    return re.sub(r"[^a-z0-9]", "", str(h).strip().lower())

def norm_email(v):
    if v is None: return ""
    s = str(v).strip().lower()
    if s in EMAIL_PLACEHOLDERS or "@" not in s or "." not in s.split("@")[-1]:
        return ""
    return s

def norm_phone(v):
    if v is None: return ""
    s = str(v).strip()
    if s.endswith(".0"): s = s[:-2]
    d = re.sub(r"\D", "", s)
    if len(d) >= 12 and d.startswith("91"): d = d[-10:]
    elif len(d) == 11 and d.startswith("0"): d = d[-10:]
    elif len(d) > 10: d = d[-10:]
    if len(d) != 10: return ""
    if d[0] not in "6789": return ""            # Indian mobile
    if len(set(d)) <= 2: return ""              # 0000000000 / 9999999999 junk
    return d

def norm_uuid(v):
    if v is None: return ""
    s = str(v).strip().lower()
    return s if UUID_RE.match(s) else ""

def parse_date(v):
    if v is None: return None
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    s = str(v).strip().split(" ")[0]
    if not s: return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%y", "%d-%m-%Y"):
        try: return dt.datetime.strptime(s, fmt).date()
        except ValueError: continue
    return None

def _has_cols(header, strict=True):
    s = {_norm(h) for h in header}
    ok = "orderno" in s and ("useremail" in s or "email" in s)
    if not strict: return ok
    return ok and ("onlydate" in s or "orderdate" in s) and ("itemstatus" in s or "laststatus" in s)

def iter_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            for row in csv.reader(f): yield row
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        firsts = {}
        for sn in wb.sheetnames:
            try: firsts[sn] = next(wb[sn].iter_rows(values_only=True))
            except StopIteration: firsts[sn] = []
        chosen = (next((sn for sn in wb.sheetnames if _has_cols(firsts[sn], True)), None)
                  or next((sn for sn in wb.sheetnames if _has_cols(firsts[sn], False)), None)
                  or wb.sheetnames[0])
        sys.stderr.write(f"[sheet] using {chosen!r} of {wb.sheetnames}\n")
        for row in wb[chosen].iter_rows(values_only=True): yield list(row)

def find_idx(header, *names):
    hl = [_norm(h) for h in header]
    for n in names:
        nl = _norm(n)
        if nl in hl: return hl.index(nl)
    return None

def main():
    path, out = sys.argv[1], sys.argv[2]
    rows = iter_rows(path)
    header = next(rows)
    i_ord = find_idx(header, "Order No.")
    i_em  = find_idx(header, "User Email", "Email")
    i_od  = find_idx(header, "Only Date", "Order Date")
    i_od2 = find_idx(header, "Order Date")
    i_st  = find_idx(header, "Item Status", "Last Status")
    i_ph  = find_idx(header, "Login Phone Number", "Mobile", "Phone")
    i_ph2 = find_idx(header, "Mobile", "Phone")
    i_uid = find_idx(header, "Username")
    assert None not in (i_ord, i_em, i_od, i_st), f"missing cols {(i_ord,i_em,i_od,i_st)} in {header}"

    orders = {}
    qc = {"file": os.path.basename(path), "lines": 0, "blank_order": 0,
          "no_email": 0, "no_phone": 0, "no_uuid": 0, "no_any_id": 0,
          "bad_date": 0, "status_counts": Counter(), "year_counts_lines": Counter()}
    maxcol = max(x for x in (i_ord,i_em,i_od,i_od2,i_st,i_ph,i_ph2,i_uid) if x is not None)

    for row in rows:
        if not row or len(row) <= maxcol: continue
        qc["lines"] += 1
        ono = row[i_ord]
        if ono is None or str(ono).strip() == "":
            qc["blank_order"] += 1; continue
        ono = str(ono).strip()
        email = norm_email(row[i_em])
        phone = norm_phone(row[i_ph])
        if not phone and i_ph2 is not None: phone = norm_phone(row[i_ph2])
        uuid = norm_uuid(row[i_uid]) if i_uid is not None else ""
        d = parse_date(row[i_od]) or (parse_date(row[i_od2]) if i_od2 is not None else None)
        st = str(row[i_st]).strip().lower() if row[i_st] is not None else ""
        qc["status_counts"][st] += 1
        if d: qc["year_counts_lines"][d.year] += 1
        else: qc["bad_date"] += 1
        paid = 0 if st in NEVER_PAID else 1
        deliv = 1 if st in FULFILLED else 0

        rec = orders.get(ono)
        if rec is None:
            orders[ono] = [d, paid, deliv, email, phone, uuid]
        else:
            if d and (rec[0] is None or d < rec[0]): rec[0] = d
            rec[1] |= paid; rec[2] |= deliv
            if not rec[3] and email: rec[3] = email
            if not rec[4] and phone: rec[4] = phone
            if not rec[5] and uuid:  rec[5] = uuid

    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["order_no","order_date","year","paid","delivered","email","phone","uuid"])
        for ono,(d,paid,deliv,email,phone,uuid) in orders.items():
            if not email: qc["no_email"] += 1
            if not phone: qc["no_phone"] += 1
            if not uuid:  qc["no_uuid"] += 1
            if not (email or phone or uuid): qc["no_any_id"] += 1
            w.writerow([ono, d.isoformat() if d else "", d.year if d else "",
                        paid, deliv, email, phone, uuid])

    qc["unique_orders"] = len(orders)
    qc["status_counts"] = dict(qc["status_counts"].most_common())
    qc["year_counts_lines"] = dict(sorted(qc["year_counts_lines"].items()))
    with open(out + ".qc.json", "w") as f: json.dump(qc, f, indent=2, default=str)
    print(json.dumps({k: qc[k] for k in
        ["file","lines","unique_orders","blank_order","no_email","no_phone",
         "no_uuid","no_any_id","bad_date","year_counts_lines"]}, default=str, indent=2))

if __name__ == "__main__":
    main()
