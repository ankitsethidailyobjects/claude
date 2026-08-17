#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os
SEG="analysis/seg"; OUT=f"{SEG}/DailyObjects_User_Modelling_Deck.pptx"
INK=RGBColor(0x0f,0x1b,0x2d); BLUE=RGBColor(0x25,0x63,0xeb); GREY=RGBColor(0x5b,0x6b,0x7f)
WHITE=RGBColor(0xff,0xff,0xff); GREEN=RGBColor(0x15,0x80,0x3d); AMBER=RGBColor(0xc2,0x56,0x0f)
LIGHT=RGBColor(0xf2,0xf5,0xf9)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
W,Hh=prs.slide_width,prs.slide_height
def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color):
    from pptx.enum.shapes import MSO_SHAPE
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid();sp.fill.fore_color.rgb=color;sp.line.fill.background();sp.shadow.inherit=False;return sp
def txt(s,x,y,w,h,text,size=18,color=INK,bold=False,align=PP_ALIGN.LEFT,italic=False,font="Calibri"):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
    p=tf.paragraphs[0];p.alignment=align;r=p.add_run();r.text=text
    r.font.size=Pt(size);r.font.bold=bold;r.font.italic=italic;r.font.color.rgb=color;r.font.name=font
    return tb
def bullets(s,x,y,w,h,items,size=15,gap=6):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(gap)
        lead,rest=(it if isinstance(it,tuple) else (None,it))
        if lead:
            r=p.add_run();r.text="• "+lead+" ";r.font.bold=True;r.font.size=Pt(size);r.font.color.rgb=INK;r.font.name="Calibri"
            r2=p.add_run();r2.text=rest;r2.font.size=Pt(size);r2.font.color.rgb=INK;r2.font.name="Calibri"
        else:
            r=p.add_run();r.text="• "+rest;r.font.size=Pt(size);r.font.color.rgb=INK;r.font.name="Calibri"
    return tb
def header(s,title,kicker=None):
    rect(s,0,0,13.333,1.15,INK)
    txt(s,0.55,0.28,11,0.7,title,26,WHITE,True)
    if kicker: txt(s,0.57,0.86,11,0.3,kicker,12,RGBColor(0x9f,0xb0,0xc4))
def pic(s,path,x,y,w=None,h=None):
    if os.path.exists(path):
        kw={}
        if w:kw["width"]=Inches(w)
        if h:kw["height"]=Inches(h)
        s.shapes.add_picture(path,Inches(x),Inches(y),**kw)
def tbl(s,x,y,w,headers,rows,colw=None,fs=11):
    from pptx.util import Inches as I
    r=len(rows)+1;c=len(headers)
    gt=s.shapes.add_table(r,c,I(x),I(y),I(w),I(0.4*r)).table
    if colw:
        for i,cw in enumerate(colw): gt.columns[i].width=I(cw)
    for j,h in enumerate(headers):
        cell=gt.cell(0,j);cell.fill.solid();cell.fill.fore_color.rgb=BLUE
        p=cell.text_frame.paragraphs[0];run=p.add_run();run.text=h;run.font.size=Pt(fs);run.font.bold=True;run.font.color.rgb=WHITE
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            cell=gt.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=WHITE if i%2 else LIGHT
            p=cell.text_frame.paragraphs[0];run=p.add_run();run.text=str(v);run.font.size=Pt(fs);run.font.color.rgb=INK
    return gt

# 1 TITLE
s=slide();rect(s,0,0,13.333,7.5,INK);rect(s,0,5.0,13.333,0.08,BLUE)
txt(s,0.8,2.1,11.7,1.4,"Customer Modelling & Retention",40,WHITE,True)
txt(s,0.8,3.2,11.7,0.8,"Cohorts · Category Affinity · Retention · MoEngage activation",22,RGBColor(0x9f,0xb0,0xc4))
txt(s,0.82,5.25,11.7,0.5,"DailyObjects · DMR 2020–Jul 2026 · 1.38M customers · 2.86M orders",14,RGBColor(0x6b,0x7d,0x93))

# 2 APPROACH
s=slide();header(s,"How we built it","Methodology")
bullets(s,0.6,1.5,12.2,5.6,[
 ("Identity resolution —","linked 1.38M unique customers across email + phone + account-UUID (fixes 2024 email masking; hub-suppressed to avoid over-merging)."),
 ("Enriched discounts —","the DMR only stores final price, hiding price-led event drops (HPS, Diwali). We detect these from weekly ASP dips — discount penetration rises 21% → 49% of orders."),
 ("Lifecycle slabs —","every customer labelled by lifetime order count: O1 (new) · O2 · 3+ (mature)."),
 ("Behavioural cohorts —","the 203K mature (3+) customers segmented via K-Means into Power Users, Deal Seekers, Light Users — on frequency, value & enriched discount reliance."),
 ("Affinity layer —","Tech vs Designer/Bags score + primary/secondary sub-category per customer."),
 ("Leakage-safe validation —","cohorts assigned at a past cut-off (31 Jul 2025) and validated on the NEXT 12 months of purchasing."),
 ("Repository —","1.38M-row master table with cohort, tiers, affinity & contact keys — ready for MoEngage."),
],size=15,gap=10)

# 3 BASE / leaky bucket
s=slide();header(s,"The base — a leaky bucket","68.6% buy once and rarely return")
pic(s,f"{SEG}/base_affinity.png",0.4,1.4,w=12.5)
txt(s,0.6,6.7,12.2,0.6,"Only 8.5% of one-time buyers return within a year; value concentrates in the ~4% who become Power Users.",13,GREY,italic=True)

# 4 DISCOUNT ENRICHMENT
s=slide();header(s,"Price-led discounts, now captured","Detected from DMR weekly-ASP dips — validated vs the elasticity sheet")
pic(s,f"{SEG}/events_2025.png",0.5,1.5,w=12.3)
txt(s,0.6,6.5,12.2,0.7,"HPS (Mar & Aug), Diwali, Black Friday and Year-end drive 30–48% price cuts in tech. Buyers who only shop these looked ‘full-price’ before — now correctly flagged as Deal Seekers.",13,GREY,italic=True)

# 5 COHORTS
s=slide();header(s,"Three cohorts — validated on future retention","All 3+ order customers; behaviour (not order-count) defines the cohort")
pic(s,f"{SEG}/cohorts_chart.png",0.35,1.45,w=12.6)
txt(s,0.6,6.55,12.2,0.7,"Power Users (28%) convert in BAU — 48% annual retention. Deal Seekers (31%) place 80% of orders on discount and lapse between sales. Light Users (42%) are low-value, mostly dormant.",13,GREY,italic=True)

# 6 RETENTION
s=slide();header(s,"Retention by cohort — monthly / quarterly / yearly","Anchor 31 Jul 2025, next 12 months")
pic(s,f"{SEG}/retention_curves.png",0.7,1.5,w=12.0)
txt(s,0.6,6.7,12.2,0.6,"Return spikes in ~Mar-2026 (next HPS) across all cohorts — repeat purchase is event-triggered, not habitual.",13,GREY,italic=True)

# 7 RETENTION TABLE
s=slide();header(s,"Return rates — the numbers","% of each cohort placing another order within…")
tbl(s,1.2,1.7,10.9,["Cohort","Share of base","1 month","1 quarter","1 year"],
 [["O1 — one-time (New)","68.6%","1.5%","3.2%","8.5%"],
  ["O2 — two orders","16.7%","3.3%","6.9%","17.4%"],
  ["Power Users","4.1%","12.4%","24.3%","48.3%"],
  ["Deal Seekers","4.5%","5.9%","12.1%","28.7%"],
  ["Light Users","6.2%","4.9%","10.3%","24.8%"]],
 colw=[3.4,2.1,1.8,1.8,1.8],fs=14)
txt(s,1.2,5.6,11,0.8,"The 1st→2nd order is the single biggest retention lever: O1 returns at 8.5% vs Power at 48%.",16,AMBER,True)

# 8 AFFINITY / MoEngage schema
s=slide();header(s,"What goes to MoEngage","Warehouse computes · MoEngage activates")
tbl(s,0.6,1.55,12.2,["MoEngage attribute","Example","Refresh"],
 [["lifecycle_cohort / order_cohort","O1 (New) / 3+ (Mature)","daily"],
  ["behavioural_cohort","Power Users / Deal Seekers / Light Users","monthly"],
  ["activity_state","Active / Lapsed / Dormant","daily"],
  ["value_tier","VIP / High / Mid / Low","weekly"],
  ["cat_affinity + tech_score","Tech (0.82)","weekly"],
  ["primary_subcategory / secondary","Wireless Charger / Designer Cases","weekly"],
  ["days_since_last · lifetime_orders · aov","47 · 6 · ₹2,600","daily"],
  ["EVENTS: cohort_changed, became_dormant, reactivated, reorder_due","—","real-time"]],
 colw=[5.2,4.6,2.4],fs=12)
txt(s,0.6,6.35,12.2,0.7,"Keyed on email/phone. Raw ML features stay in the warehouse; MoEngage receives finished decisions only.",13,GREY,italic=True)

# 9 GAPS + STRATEGY
s=slide();header(s,"Gaps & the retention strategy")
txt(s,0.5,1.3,6.2,0.4,"WHERE WE LOSE USERS",14,AMBER,True)
bullets(s,0.5,1.75,6.3,5.2,[
 ("Leaky top —","68.6% one-time, 8.5% return."),
 ("O2→O3 stall —","17% before the loyalty inflection."),
 ("Discount addiction —","31% Deal Seekers, 80% on discount."),
 ("Event-only repeat —","little BAU habit."),
 ("Category concentration —","74% Tech, bags under-sold."),
 ("Dormancy —","59% of mature dormant."),
],size=14,gap=8)
txt(s,6.95,1.3,6,0.4,"WHAT WE DO",14,GREEN,True)
bullets(s,6.95,1.75,5.9,5.2,[
 ("P1 Win the 2nd order —","onboarding + first-repeat incentive."),
 ("P2 Push to 3rd —","cross-category + reorder timing."),
 ("P3 Grow Power —","VIP, full-price Tech→Bags cross-sell."),
 ("P4 De-risk Deal Seekers —","bundles, raise AOV, wean off discount."),
 ("P5 Kill event-dependence —","always-on replenishment journeys."),
 ("P6 Reactivate / sunset —","sized by value & affinity."),
],size=14,gap=8)

# 10 REFRESH
s=slide();header(s,"Operating the system","Refresh cadence")
tbl(s,1.0,1.6,11.3,["Layer","Refresh","Note"],
 [["Lifecycle cohort, activity, recency","Daily","Deterministic / fast-moving"],
  ["Category & sub-category affinity","Daily–Weekly","Drives creative"],
  ["Value tier, discount class","Weekly","Slow ratios"],
  ["Behavioural cohort (re-score)","Monthly","Hysteresis-guarded"],
  ["Model + event windows (re-train)","Quarterly","Identity-changing"]],
 colw=[5.0,2.8,3.5],fs=13)
txt(s,1.0,5.3,11,1.4,"Analysts refresh via Claude in plain language: “Refresh the cohort model with the latest DMR, rebuild the repository, and push to MoEngage.” Claude runs the pipeline, sanity-checks cohort sizes, and syncs.",15,INK,False)

prs.save(OUT); print("saved",OUT,"slides:",len(prs.slides._sldIdLst))
