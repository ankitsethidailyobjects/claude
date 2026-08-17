#!/usr/bin/env python3
"""Lean extractor: per (order_no, sub_category) -> gmv, qty. Falls back to
top-level Category where Sub Category is absent (older files)."""
import sys, os, csv, re
from collections import defaultdict
def _norm(h): return re.sub(r"[^a-z0-9]","",str(h).strip().lower()) if h is not None else ""
def num(v):
    if v is None: return 0.0
    try: return float(str(v).replace(",","").strip() or 0)
    except: return 0.0
def has_cols(header):
    s={_norm(h) for h in header}; return "orderno" in s and ("subcategory" in s or "category" in s)
def rows(path):
    ext=os.path.splitext(path)[1].lower()
    if ext==".csv":
        with open(path,newline="",encoding="utf-8-sig",errors="replace") as f:
            for r in csv.reader(f): yield r
    else:
        import openpyxl
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True); firsts={}
        for sn in wb.sheetnames:
            try: firsts[sn]=next(wb[sn].iter_rows(values_only=True))
            except StopIteration: firsts[sn]=[]
        ch=next((sn for sn in wb.sheetnames if has_cols(firsts[sn])),wb.sheetnames[0])
        for r in wb[ch].iter_rows(values_only=True): yield list(r)
def idx(header,*names):
    hl=[_norm(h) for h in header]
    for n in names:
        if _norm(n) in hl: return hl.index(_norm(n))
    return None
def main():
    path,out=sys.argv[1],sys.argv[2]
    it=rows(path); header=next(it)
    i_ord=idx(header,"Order No."); i_sub=idx(header,"Sub Category"); i_cat=idx(header,"Category")
    i_gmv=idx(header,"Product Grand Total(FF)","Product Grand Total","Product final Price"); i_qty=idx(header,"Qty Ordered")
    subcol=i_sub if i_sub is not None else i_cat
    assert i_ord is not None and subcol is not None
    maxc=max(x for x in (i_ord,subcol,i_gmv,i_qty) if x is not None)
    agg=defaultdict(lambda:[0.0,0.0])
    for r in it:
        if not r or len(r)<=maxc: continue
        o=r[i_ord]
        if o is None or str(o).strip()=="": continue
        o=str(o).strip()
        sub=str(r[subcol]).strip() if r[subcol] is not None and str(r[subcol]).strip() else "Unknown"
        gmv=num(r[i_gmv]) if i_gmv is not None else 0.0
        qty=num(r[i_qty]) if i_qty is not None else 1.0
        a=agg[(o,sub)]; a[0]+=gmv; a[1]+=qty
    with open(out,"w",newline="") as f:
        w=csv.writer(f); w.writerow(["order_no","sub_category","gmv","qty"])
        for (o,sub),(gmv,qty) in agg.items(): w.writerow([o,sub,round(gmv,2),int(qty)])
    print(f"{os.path.basename(path)}: rows={len(agg)}")
if __name__=="__main__": main()
