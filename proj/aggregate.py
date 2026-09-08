#!/usr/bin/env python3
"""Aggregate one DMR file -> per-SKU and category rollups. Streams; low memory.
Usage: python3 aggregate.py <path> <label>   e.g. 2026-07"""
import sys, os, csv, json
import pandas as pd

EXCLUDE = {"paymentfailed","paymentpending","cancel"}  # not real fulfillable demand

WANT = ["Item Status","Qty Ordered","Only Date","SKU","Category","Sub Category","Selling Price"]
# canonical key = lowercase, strip, spaces/underscores removed
def canon(s): return str(s).strip().lower().replace(" ","").replace("_","")
CANON2WANT = {canon(w): w for w in WANT}

def norm(s):
    return (s or "").strip()

def process_csv(path):
    # accumulate per SKU
    acc = {}
    for chunk in pd.read_csv(path, usecols=lambda c: canon(c) in CANON2WANT, dtype=str,
                             low_memory=True, chunksize=100000):
        chunk.columns = [CANON2WANT.get(canon(c), c.strip()) for c in chunk.columns]
        st = chunk["Item Status"].astype(str).str.strip().str.lower()
        keep = ~st.isin(EXCLUDE)
        c = chunk[keep].copy()
        c["qty"] = pd.to_numeric(c["Qty Ordered"].astype(str).str.replace(',',''), errors='coerce').fillna(0)
        c["price"] = pd.to_numeric(c["Selling Price"].astype(str).str.replace(',',''), errors='coerce').fillna(0)
        c["gross"] = c["qty"]*c["price"]
        c["deliv"] = ((st[keep]=="delivered")*c["qty"]).values
        c["day"] = c["Only Date"].astype(str).str.strip()
        for sku, grp in c.groupby("SKU"):
            a = acc.get(sku)
            if a is None:
                a = acc[sku] = {"qty":0.0,"gross":0.0,"deliv":0.0,"days":set(),
                                "cat":{}, "sub":{}, "pxnum":0.0,"pxden":0.0}
            a["qty"]+=grp["qty"].sum(); a["gross"]+=grp["gross"].sum(); a["deliv"]+=grp["deliv"].sum()
            a["days"].update(grp["day"].unique())
            # modal category by qty
            for cat,q in grp.groupby("Category")["qty"].sum().items():
                a["cat"][cat]=a["cat"].get(cat,0)+q
            for sc,q in grp.groupby("Sub Category")["qty"].sum().items():
                a["sub"][sc]=a["sub"].get(sc,0)+q
            a["pxnum"]+=grp["gross"].sum(); a["pxden"]+=grp["qty"].sum()
    return acc

def process_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # pick sheet whose header row contains the most WANT columns (skip pivot sheets)
    best=None
    for sn in wb.sheetnames:
        ws=wb[sn]
        try:
            first=next(ws.iter_rows(values_only=True, max_row=1))
        except StopIteration:
            continue
        canons={canon(x) for x in first if x is not None}
        score=len(canons & set(CANON2WANT))
        if best is None or score>best[0]:
            best=(score, sn)
    ws = wb[best[1]]
    print(f"  [xlsx] using sheet '{best[1]}' (matched {best[0]}/{len(WANT)} cols)")
    it = ws.iter_rows(values_only=True)
    header = [CANON2WANT.get(canon(x), norm(str(x))) for x in next(it)]
    idx = {name: (header.index(name) if name in header else None) for name in WANT}
    acc = {}
    for row in it:
        st = norm(str(row[idx["Item Status"]])).lower() if idx["Item Status"] is not None else ""
        if st in EXCLUDE: continue
        sku = norm(str(row[idx["SKU"]])) if idx["SKU"] is not None else ""
        if not sku or sku=="None": continue
        try: qty=float(str(row[idx["Qty Ordered"]]).replace(',','')) if idx["Qty Ordered"] is not None and row[idx["Qty Ordered"]] not in (None,'') else 0.0
        except: qty=0.0
        try: px=float(str(row[idx["Selling Price"]]).replace(',','')) if idx["Selling Price"] is not None and row[idx["Selling Price"]] not in (None,'') else 0.0
        except: px=0.0
        gross=qty*px
        cat=norm(str(row[idx["Category"]])) if idx["Category"] is not None else ""
        sub=norm(str(row[idx["Sub Category"]])) if idx["Sub Category"] is not None else ""
        day=norm(str(row[idx["Only Date"]])) if idx["Only Date"] is not None else ""
        a=acc.get(sku)
        if a is None:
            a=acc[sku]={"qty":0.0,"gross":0.0,"deliv":0.0,"days":set(),"cat":{},"sub":{},"pxnum":0.0,"pxden":0.0}
        a["qty"]+=qty; a["gross"]+=gross; a["deliv"]+= qty if st=="delivered" else 0.0
        if day: a["days"].add(day)
        if cat: a["cat"][cat]=a["cat"].get(cat,0)+qty
        if sub: a["sub"][sub]=a["sub"].get(sub,0)+qty
        a["pxnum"]+=gross; a["pxden"]+=qty
    wb.close()
    return acc

def main():
    path, label = sys.argv[1], sys.argv[2]
    acc = process_xlsx(path) if path.lower().endswith(('.xlsx','.xls')) else process_csv(path)
    rows=[]
    for sku,a in acc.items():
        cat = max(a["cat"], key=a["cat"].get) if a["cat"] else ""
        sub = max(a["sub"], key=a["sub"].get) if a["sub"] else ""
        asp = (a["pxnum"]/a["pxden"]) if a["pxden"] else 0.0
        rows.append({"label":label,"SKU":sku,"category":cat,"sub_category":sub,
                     "demand_units":round(a["qty"],2),"gross_rev":round(a["gross"],2),
                     "delivered_units":round(a["deliv"],2),"days_sold":len(a["days"]),"asp":round(asp,2)})
    df=pd.DataFrame(rows).sort_values("gross_rev",ascending=False)
    out=f"agg/skus_{label}.csv"; df.to_csv(out,index=False)
    # category rollup
    cm=df.groupby("category").agg(demand_units=("demand_units","sum"),gross_rev=("gross_rev","sum"),skus=("SKU","nunique")).reset_index()
    cm.insert(0,"label",label)
    cmpath="agg/cat_month.csv"
    cm.to_csv(cmpath, mode='a', header=not os.path.exists(cmpath), index=False)
    tot=df["gross_rev"].sum(); u=df["demand_units"].sum()
    print(f"{label}: SKUs={len(df)}  units={u:,.0f}  gross=Rs{tot:,.0f} ({tot/1e7:,.2f} cr)  -> {out}")

if __name__=="__main__":
    main()
