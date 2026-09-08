#!/usr/bin/env python3
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

d=pd.read_csv('agg/parent_compare_full.csv')
top=d.sort_values('rank_val',ascending=False).head(100).reset_index(drop=True)

SH_OCT={'Cases And Sleeves':1.40,'Accessories':0.83,'Bags and Backpacks':1.09,'Home Office':0.95,
        'Wallet and Pouches':0.55,'Gifts':1.65,'Travel':0.75}
SH_NOV={'Cases And Sleeves':1.21,'Accessories':0.84,'Bags and Backpacks':1.24,'Home Office':1.07,
        'Wallet and Pouches':0.55,'Gifts':0.82,'Travel':0.87}

def reason(r):
    cat=r['cat']; base=r['base_u']; fo=SH_OCT.get(cat,1.0); fn=SH_NOV.get(cat,1.0)
    parts=[f"Base {base:.0f} u/mo (Jul+Aug'26 run-rate); {cat or 'unmapped'} festive x{fo:.2f} Oct / x{fn:.2f} Nov."]
    to,mo=r['t_oct_u'],r['m_oct_u']
    if to==0:
        parts.append(f"TEAM OMITTED (planned 0) — we carry {mo:.0f} u from live run-rate.")
    elif base>0 and to>1.6*max(mo,1):
        parts.append(f"Team OVER-planned: their Oct {to:.0f} u = {to/base:.1f}x run-rate; we anchor to run-rate ({mo:.0f} u).")
    elif base>0 and to<0.6*mo:
        parts.append(f"Team UNDER-planned: their Oct {to:.0f} u vs our run-rate-based {mo:.0f} u.")
    else:
        parts.append(f"In line with team (their Oct {to:.0f} vs our {mo:.0f} u).")
    if cat=='Bags and Backpacks':
        parts.append("Nov>Oct (bags peak in Nov'25); team used flat -5%.")
    elif cat=='Cases And Sleeves':
        parts.append("Cases ease Oct->Nov (gifting front-loads to Oct); team used flat -5%.")
    if r['base_g']>0 and r['hps_g']>3*r['base_g']:
        parts.append(f"Event-led (Sep-HPS {r['hps_g']/max(r['base_g'],1):.0f}x base) — held at base, no HPS in Oct/Nov.")
    return " ".join(parts)

top['Reasoning']=top.apply(reason,axis=1)
out=pd.DataFrame({
 'Rank':range(1,len(top)+1),
 'Parent':top['parent'],'Category':top['cat'],
 'Child SKUs (ours)':top['child'].astype(int),
 'Team Oct u':top['t_oct_u'].round(0),'Our Oct u':top['m_oct_u'].round(0),'Δ Oct u':top['d_oct_u'].round(0),
 'Team Nov u':top['t_nov_u'].round(0),'Our Nov u':top['m_nov_u'].round(0),'Δ Nov u':top['d_nov_u'].round(0),
 'Team Oct ₹':top['t_oct_rev'].round(0),'Our Oct ₹':top['m_oct_rev'].round(0),
 'Reasoning / methodology':top['Reasoning']})

# ---- write single-sheet workbook ----
NAVY="1F3864"; BLUE="2E5496"; GREY="F2F2F2"; YEL="FFF2CC"; GREEN="C6EFCE"; RED="F8CBAD"
thin=Side(style='thin',color="BFBFBF"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="Top100 Parent Compare"
ws.sheet_view.showGridLines=False
# title
ncol=len(out.columns)
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncol)
c=ws.cell(1,1,"Top 100 Parent SKUs — Team Estimate vs Inventory Demand (Oct/Nov 2026), with methodology")
c.font=Font(bold=True,size=13,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.alignment=Alignment(vertical='center')
ws.row_dimensions[1].height=24
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol)
c=ws.cell(2,1,"Units = inventory demand. Our numbers: Jul+Aug'26 run-rate → 2025 seasonal category shift → reconciled to ₹12cr/₹11.5cr. Δ = Ours − Team (green: we plan more; red: we plan less). Team omissions shown where Team = 0.")
c.font=Font(size=9,italic=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=BLUE); c.alignment=Alignment(vertical='center',wrap_text=True); ws.row_dimensions[2].height=28
# header
hr=3
for j,name in enumerate(out.columns,1):
    cc=ws.cell(hr,j,name); cc.font=Font(bold=True,color="FFFFFF",size=9)
    cc.fill=PatternFill("solid",fgColor=BLUE); cc.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cc.border=border
moneycols={'Team Oct ₹','Our Oct ₹'}; intcols={'Child SKUs (ours)','Team Oct u','Our Oct u','Δ Oct u','Team Nov u','Our Nov u','Δ Nov u'}
for i,(_,row) in enumerate(out.iterrows(),hr+1):
    for j,name in enumerate(out.columns,1):
        v=row[name]; cc=ws.cell(i,j,v); cc.border=border; cc.font=Font(size=9)
        if name in moneycols: cc.number_format='#,##0'
        elif name in intcols: cc.number_format='#,##0'
        if name=='Reasoning / methodology': cc.alignment=Alignment(wrap_text=True,vertical='top')
        elif name in ('Parent','Category'): cc.alignment=Alignment(vertical='top')
        else: cc.alignment=Alignment(horizontal='center',vertical='top')
        if i%2==0 and name!='Reasoning / methodology': cc.fill=PatternFill("solid",fgColor=GREY)
    ws.row_dimensions[i].height=30
# conditional color on delta columns
last=hr+len(out)
for col in ['Δ Oct u','Δ Nov u']:
    L=get_column_letter(list(out.columns).index(col)+1); rng=f"{L}{hr+1}:{L}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill("solid",fgColor=GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill("solid",fgColor=RED)))
# widths
W=[5,30,19,9,9,9,8,9,9,8,12,12,72]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='C4'
# totals row
tr=last+1
ws.cell(tr,2,"TOP-100 TOTAL").font=Font(bold=True)
for name in ['Team Oct u','Our Oct u','Δ Oct u','Team Nov u','Our Nov u','Δ Nov u','Team Oct ₹','Our Oct ₹']:
    j=list(out.columns).index(name)+1; cc=ws.cell(tr,j,out[name].sum())
    cc.number_format='#,##0'; cc.font=Font(bold=True); cc.fill=PatternFill("solid",fgColor=YEL); cc.border=border
wb.save('out/Top100_Parent_Team_vs_Ours_OctNov2026.xlsx')
print("saved out/Top100_Parent_Team_vs_Ours_OctNov2026.xlsx")
print("omitted-by-team in top100:",(out['Team Oct u']==0).sum())
print("we plan MORE (ΔOct>0):",(out['Δ Oct u']>0).sum()," we plan LESS:",(out['Δ Oct u']<0).sum())
print(out[['Parent','Team Oct u','Our Oct u','Δ Oct u']].head(12).to_string(index=False))
